"""Synthetic regressions for teacher selections, document objects and lesson tasks."""
import base64
import datetime
import io
import json
import shutil
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch

from django.test import TestCase, SimpleTestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone

from .ai_context import build_assessment_request, evidence_cache, extract_file_evidence
from .models import AIJob, Assignment, AssignmentScheduleTarget, BellSchedule, Subject, TeacherLessonSchedule, SubmissionFile
from .office_objects import docx_object_evidence, ooxml_object_evidence
from . import tests_ai_assessment_v4 as v4_tests
from .utils import get_published_tasks_for_lesson, get_teacher_live_lesson_status, get_teacher_upcoming_notifications


@override_settings(AI_JOBS_EAGER=False)
class TeacherSelectionAndScheduleTests(TestCase):
    setUp = v4_tests.AssessmentV4Tests.setUp
    result = v4_tests.AssessmentV4Tests.result
    evaluate = v4_tests.AssessmentV4Tests.evaluate

    def test_selected_definitions_reach_provider_and_unselected_results_are_removed(self):
        result, call = self.evaluate(self.result(gr_results=[{'code': 'ГР 1', 'grade': 8}, {'code': 'ГР 10', 'grade': 2}]))
        prompt = call.call_args.kwargs['prompt_text']
        context = next(json.loads(line) for line in prompt.splitlines() if line.startswith('{"class"'))
        self.assertEqual(context['active_result_groups'], ['ГР 1'])
        self.assertEqual(context['result_group_definitions'], [{'code': 'ГР 1', 'name': 'Зміст'}])
        self.assertEqual([row['code'] for row in result['gr_results']], ['ГР 1'])

    def test_both_endpoints_preserve_checkbox_choice_and_freeze_defaults(self):
        self.client.force_login(self.user)
        for route in ('ai_check_single_submission', 'api_ai_process_item'):
            for fields, expected in (({}, ['ГР 1']),
                                     ({'preset_id': self.preset.pk, 'selected_gr_codes': '["ГР 10"]'}, ['ГР 10']),
                                     ({'selected_gr_codes': ['ГР 1', 'ГР 10']}, ['ГР 1', 'ГР 10'])):
                with self.subTest(route=route, fields=fields):
                    AIJob.objects.all().delete()
                    response = self.client.post(reverse(route, args=[self.sub.pk]), fields)
                    self.assertEqual(response.status_code, 202, response.content)
                    job = AIJob.objects.get(submission=self.sub)
                    self.assertEqual(job.parameters['preset_id'], str(self.preset.pk))
                    self.assertEqual(job.parameters['selected_gr_codes'], expected)

    def test_invalid_or_empty_selection_never_enqueues_on_either_endpoint(self):
        self.client.force_login(self.user)
        for route in ('ai_check_single_submission', 'api_ai_process_item'):
            for fields in ({'selected_gr_codes': '[]'}, {'selected_gr_codes': '["ГР 2"]'},
                           {'selected_gr_codes': '"ГР 1"'}, {'selected_gr_codes': '[1]'},
                           {'preset_id': 'not-an-id'}, {'preset_id': 999999}):
                response = self.client.post(reverse(route, args=[self.sub.pk]), fields)
                self.assertEqual(response.status_code, 400, (route, fields, response.content))
        self.assertFalse(AIJob.objects.exists())

    def test_preset_without_group_definitions_cannot_silently_ignore_selected_codes(self):
        self.traditional.evaluation_type = 'nus'
        self.traditional.save()
        self.client.force_login(self.user)
        for route in ('ai_check_single_submission', 'api_ai_process_item'):
            response = self.client.post(reverse(route, args=[self.sub.pk]),
                            {'preset_id': self.traditional.pk, 'selected_gr_codes': '["ГР 1"]'})
            self.assertEqual(response.status_code, 400)
            self.assertIn('немає визначених ГР', response.json()['error'])
        self.assertFalse(AIJob.objects.exists())

    def test_primary_file_and_extra_attachments_are_read_once_and_missing_files_reported(self):
        from .gemini_service import extract_submission_content
        extra = SubmissionFile.objects.create(submission=self.sub, original_name='extra.txt',
                        file=SimpleUploadedFile('extra.txt', 'Додаткова відповідь'.encode()))
        text, _, error = extract_submission_content(self.sub)
        self.assertIsNone(error)
        self.assertIn('Безпечні паролі.', '\n'.join(text))
        self.assertIn('Додаткова відповідь', '\n'.join(text))
        SubmissionFile.objects.create(submission=self.sub, original_name='mirror.txt', file=self.sub.file.name)
        text, _, _ = extract_submission_content(self.sub)
        self.assertEqual('\n'.join(text).count('Безпечні паролі.'), 1)
        Path(extra.file.path).unlink()
        text, _, _ = extract_submission_content(self.sub)
        self.assertTrue(any('extra.txt' in item and 'недоступний на сервері' in item for item in text))

    def prepare_schedule(self):
        self.now = timezone.make_aware(datetime.datetime(2026, 10, 5, 9, 50))
        self.subject = Subject.objects.create(name='Інформатика')
        self.teacher.subjects.add(self.subject)
        self.slot = BellSchedule.objects.create(lesson_number=2, start_time=datetime.time(10), end_time=datetime.time(10, 45))
        self.other_slot = BellSchedule.objects.create(lesson_number=3, start_time=datetime.time(11), end_time=datetime.time(11, 45))
        self.lesson = TeacherLessonSchedule.objects.create(teacher=self.teacher, class_group=self.group,
                    subject=self.subject, day_of_week=1, bell_slot=self.slot)
        self.assignment.subject = self.subject
        self.assignment.published_at = self.now
        self.assignment.save()
        self.target = AssignmentScheduleTarget.objects.create(assignment=self.assignment, class_group=self.group,
                    target_date=self.now.date(), target_day_of_week=1, bell_slot=self.slot)

    def test_widget_and_notifications_agree_on_exact_lesson(self):
        self.prepare_schedule()
        status = get_teacher_live_lesson_status(self.teacher, self.now)
        task = status['next_lesson_task']
        self.assertTrue(task['has_task'])
        self.assertEqual(task['url'], reverse('assignment_detail', args=[self.assignment.pk]))
        self.assertFalse(any(n['type'] == 'missing_task' for n in get_teacher_upcoming_notifications(self.teacher, self.now)))
        self.target.bell_slot = self.other_slot
        self.target.save()
        self.assertFalse(get_teacher_live_lesson_status(self.teacher, self.now)['next_lesson_task']['has_task'])
        self.assertTrue(any(n['type'] == 'missing_task' for n in get_teacher_upcoming_notifications(self.teacher, self.now)))

    def test_other_dates_subjects_drafts_and_individual_tasks_do_not_count(self):
        self.prepare_schedule()
        for changes in ({'subject': Subject.objects.create(name='Математика')}, {'status': 'draft'}, {'is_individual': True}):
            for key, value in changes.items():
                setattr(self.assignment, key, value)
            self.assignment.save()
            self.assertEqual(get_published_tasks_for_lesson(self.teacher, self.lesson, self.now.date()), [])
            self.assignment.subject, self.assignment.status, self.assignment.is_individual = self.subject, 'published', False
        self.assignment.save()
        self.target.target_date += datetime.timedelta(days=7)
        self.target.save()
        self.assertEqual(get_published_tasks_for_lesson(self.teacher, self.lesson, self.now.date()), [])
        self.target.target_date = None
        self.target.save()
        self.assignment.published_at -= datetime.timedelta(days=7)
        self.assignment.save()
        self.assertEqual(get_published_tasks_for_lesson(self.teacher, self.lesson, self.now.date()), [])

    def test_next_week_and_live_api_use_same_task_state(self):
        self.prepare_schedule()
        self.now = self.now.replace(day=4)  # Sunday -> Monday 5 October.
        status = get_teacher_live_lesson_status(self.teacher, self.now)
        self.assertEqual(status['next_lesson_task']['lesson_date'], '2026-10-05')
        self.assertTrue(status['next_lesson_task']['has_task'])
        self.client.force_login(self.user)
        with patch('feed.views.timezone.now', return_value=self.now):
            live = self.client.get(reverse('teacher_live_status')).json()['live_status']
        self.assertEqual(live['next_lesson_task'], status['next_lesson_task'])

    def test_create_link_prefills_class_subject_date_and_slot(self):
        self.prepare_schedule()
        self.assignment.status = 'draft'
        self.assignment.save()
        task = get_teacher_live_lesson_status(self.teacher, self.now)['next_lesson_task']
        self.client.force_login(self.user)
        response = self.client.get(task['url'])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['form'].initial['classes'], [self.group.pk])
        self.assertEqual(response.context['form'].initial['subject'], self.subject.pk)
        self.assertEqual(response.context['existing_target_dates'], {self.group.pk: '2026-10-05'})
        self.assertTrue(any(getattr(s, 'is_selected_for_assignment', False) for s in response.context['teacher_lesson_schedules']))
        result = self.client.post(reverse('assignment_create'), {
            'title': 'Нове завдання для уроку', 'description': 'Створіть карту знань.',
            'classes': [self.group.pk], 'subject': self.subject.pk, 'publish_choice': 'now',
            f'class_schedule_target_{self.group.pk}': f'slot_1_{self.slot.pk}',
            f'class_schedule_target_date_{self.group.pk}': '2026-10-05',
        })
        self.assertEqual(result.status_code, 302)
        self.assertTrue(get_teacher_live_lesson_status(self.teacher, self.now)['next_lesson_task']['has_task'])


class DocxObjectsTests(SimpleTestCase):
    def setUp(self):
        evidence_cache().clear()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name, 'mindmap.docx')

    def make_document(self, extra_objects=True):
        from docx import Document
        from PIL import Image
        document = Document()
        document.add_paragraph('Пояснення учня')
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Дані таблиці'
        table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = 'Вкладена таблиця'
        document.sections[0].header.paragraphs[0].text = 'Заголовок'
        image = io.BytesIO()
        Image.new('RGB', (80, 40), 'green').save(image, format='PNG')
        image.seek(0)
        document.add_picture(image)
        document.save(self.path)
        if not extra_objects:
            return
        with zipfile.ZipFile(self.path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        body = ('<w:p><w:r><w:drawing xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                '<a:t>Текст у фігурі карти</a:t></w:drawing></w:r></w:p>')
        parts['word/document.xml'] = parts['word/document.xml'].replace(b'</w:body>', (body + '</w:body>').encode())
        parts['word/diagrams/data1.xml'] = ('''<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"
          xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst>
          <dgm:pt modelId="root"><dgm:t><a:p><a:r><a:t>Карта знань</a:t></a:r></a:p></dgm:t></dgm:pt>
          <dgm:pt modelId="branch"><dgm:t><a:p><a:r><a:t>Безпека паролів</a:t></a:r></a:p></dgm:t></dgm:pt>
          </dgm:ptLst><dgm:cxnLst><dgm:cxn srcId="root" destId="branch" type="parOf"/></dgm:cxnLst></dgm:dataModel>''').encode()
        parts['word/embeddings/oleObject1.bin'] = b'test-only-placeholder'
        with zipfile.ZipFile(self.path, 'w') as archive:
            for name, data in parts.items():
                archive.writestr(name, data)

    def test_smartart_shape_text_edges_and_embedded_files_are_visible(self):
        self.make_document()
        with patch('feed.ai_context._office_pdf', side_effect=ValueError('conversion unavailable')):
            evidence = extract_file_evidence(str(self.path))
        for text in ('Пояснення учня', 'Дані таблиці', 'Вкладена таблиця', 'Заголовок', 'Карта знань', 'Безпека паролів', 'Текст у фігурі карти', '"srcId": "root"', '"destId": "branch"'):
            self.assertIn(text, evidence['text'])
        self.assertTrue(any(item['mime_type'] == 'image/png' for item in evidence['media']))
        self.assertTrue(any('Візуальний вигляд' in item for item in evidence['limitations']))
        self.assertTrue(any('вкладені файли' in item for item in evidence['limitations']))

    def test_original_image_is_sent_even_if_pdf_conversion_succeeds(self):
        self.make_document()
        with patch('feed.ai_context._office_pdf', return_value=b'%PDF-test'):
            evidence = extract_file_evidence(str(self.path))
        self.assertEqual([item['mime_type'] for item in evidence['media']], ['application/pdf', 'image/png'])
        self.assertNotIn('Візуальний вигляд не прочитано', ' '.join(evidence['limitations']))

    def test_graphic_docx_preview_shows_pages_and_retains_copyable_text(self):
        from .utils import convert_docx_to_html
        from .review_preview import ASSET_MARKER, assets_available
        from PIL import Image
        self.make_document()
        image = io.BytesIO()
        Image.new('RGB', (80, 100), 'green').save(image, format='JPEG')
        pages = [{'mime_type': 'image/jpeg', 'data': base64.b64encode(image.getvalue()).decode()}]
        with patch('feed.ai_context._office_pdf', return_value=b'%PDF-test'), \
             patch('feed.ai_context.media_for_provider', return_value=pages):
            content, error = convert_docx_to_html(str(self.path), preview_assets=True)
        self.assertIsNone(error)
        self.assertIn('Сторінка 1 документа зі схемами', content)
        self.assertIn('Текст документа для читання та копіювання', content)
        self.assertIn('Пояснення учня', content)
        self.assertIn(ASSET_MARKER, content)
        self.assertTrue(assets_available(str(self.path), content))

    def test_failed_graphic_preview_keeps_text_and_reports_missing_visual_objects(self):
        from .utils import convert_docx_to_html
        self.make_document()
        with patch('feed.ai_context._office_pdf', side_effect=ValueError('unavailable')):
            content, error = convert_docx_to_html(str(self.path), preview_assets=True)
        self.assertIn('Пояснення учня', content)
        self.assertIn('Схеми та фігури', error)
        self.assertIn('відкрийте оригінальний файл', error)

    def test_large_raster_and_vector_limitations_are_not_silently_skipped(self):
        from PIL import Image
        self.make_document()
        image = io.BytesIO()
        Image.new('RGB', (1800, 1800), 'blue').save(image, format='BMP')
        self.assertGreater(len(image.getvalue()), 8 * 1024 * 1024)
        with zipfile.ZipFile(self.path, 'a') as archive:
            archive.writestr('word/media/large.bmp', image.getvalue())
            archive.writestr('word/media/vector.emf', b'unsupported-vector-fixture')
        with patch('feed.ai_context._office_pdf', side_effect=ValueError('conversion unavailable')):
            evidence = extract_file_evidence(str(self.path))
        self.assertEqual(len(evidence['media']), 2)
        self.assertTrue(any('large.bmp' in item['source'] for item in evidence['media']))
        self.assertTrue(any('vector.emf' in item for item in evidence['limitations']))

    def test_presentation_and_spreadsheet_diagram_parts_are_read_too(self):
        self.make_document()
        with zipfile.ZipFile(self.path) as archive:
            graph = archive.read('word/diagrams/data1.xml')
        for prefix in ('ppt', 'xl'):
            with zipfile.ZipFile(self.path, 'w') as archive:
                archive.writestr(prefix + '/diagrams/data1.xml', graph)
            evidence = ooxml_object_evidence(self.path, prefix)
            self.assertIn('Безпека паролів', '\n'.join(evidence['text']))

    def test_unreadable_diagram_is_explicitly_reported_without_losing_other_parts(self):
        self.make_document()
        with zipfile.ZipFile(self.path, 'a') as archive:
            archive.writestr('word/diagrams/data2.xml', b'<broken')
        evidence = docx_object_evidence(self.path)
        self.assertIn('Безпека паролів', '\n'.join(evidence['text']))
        self.assertTrue(any('data2.xml' in item for item in evidence['limitations']))

    def test_real_libreoffice_conversion_retains_text_and_visual_media(self):
        if not shutil.which('libreoffice'):
            self.skipTest('LibreOffice is unavailable')
        self.make_document(extra_objects=False)
        evidence = extract_file_evidence(str(self.path))
        self.assertEqual([item['mime_type'] for item in evidence['media']], ['application/pdf', 'image/png'])
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(base64.b64decode(evidence['media'][0]['data'])))
        self.assertGreaterEqual(len(reader.pages), 1)
        self.assertIn('Пояснення учня', '\n'.join(page.extract_text() for page in reader.pages))
