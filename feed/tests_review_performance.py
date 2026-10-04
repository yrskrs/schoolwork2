"""Review optimizations preserve permissions, pixels, queue semantics and grading."""
import io
import re
import os
from concurrent.futures import ThreadPoolExecutor
from unittest import skipUnless
from pathlib import Path
from unittest.mock import patch
from urllib.parse import urlsplit, parse_qs

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from lxml import html
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from .middleware import set_has_admin
from .models import Assignment, AssignmentFile, ClassGroup, School, Submission, SubmissionFile, Teacher
from .review_preview import asset_directory, convert_office


class ReviewPerformanceTests(TestCase):
    def setUp(self):
        set_has_admin(True)
        self.user = User.objects.create_user('fast-review-owner', password='test-password')
        self.teacher = Teacher.objects.create(user=self.user, full_name='Вчитель')
        School.objects.create(name='Тестова школа', admin=self.user)
        self.group = ClassGroup.objects.create(name='7-А', grade=7)
        self.teacher.classes.add(self.group)
        self.assignment = Assignment.objects.create(teacher=self.teacher, title='Презентація', status='published')
        self.assignment.classes.add(self.group)
        image = io.BytesIO()
        Image.new('RGB', (80, 60), (10, 20, 30)).save(image, 'PNG')
        self.image_bytes = image.getvalue()
        presentation = Presentation()
        for n in range(3):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(io.BytesIO(self.image_bytes), Inches(1), Inches(1))
            slide.shapes.add_textbox(Inches(1), Inches(3), Inches(5), Inches(1)).text = f'Відповідь {n}'
        stream = io.BytesIO(); presentation.save(stream)
        self.sub = Submission.objects.create(assignment=self.assignment, teacher=self.teacher,
            class_group=self.group, first_name='Учень', last_name='Перший',
            file=SimpleUploadedFile('answer.pptx', stream.getvalue()))
        self.client.force_login(self.user)
        self.view_url = reverse('view_file', args=[self.sub.pk])
        self.preview_url = reverse('review_document_preview', args=[self.sub.pk])
        self.duplicates_url = reverse('review_duplicates', args=[self.sub.pk])

    def preview(self, **params):
        response = self.client.get(self.preview_url, params)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        return html.fromstring(response.json()['html'])

    def test_shell_does_not_convert_or_scan_files_and_never_claims_a_clean_result(self):
        with (patch('feed.views.convert_pptx_to_html', side_effect=AssertionError('blocking conversion')),
              patch('feed.views.check_submission_duplicates', side_effect=AssertionError('blocking scan'))):
            response = self.client.get(self.view_url)
        self.assertContains(response, 'data-review-preview=')
        self.assertContains(response, 'data-review-duplicates=')
        page = html.fromstring(response.content)
        chip = page.get_element_by_id('fv-quick-duplicate')
        self.assertIn('Перевіряємо', chip.text_content())
        self.assertNotIn('не виявлено', chip.text_content())
        self.assertIn('review_sync=1', response.context['review_sync_url'])

    def test_preview_keeps_all_slides_and_serves_identical_images_once(self):
        page = self.preview()
        self.assertEqual(len(page.xpath('//*[contains(@class,"pptx-slide-card")]')), 3)
        sources = page.xpath('//img/@src')
        self.assertEqual(len(sources), 3)
        self.assertEqual(len(set(sources)), 1)
        self.assertNotIn('data:image/', html.tostring(page).decode())
        self.assertEqual(page.xpath('//img/@loading'), ['lazy'] * 3)
        self.assertEqual(page.xpath('//img/@decoding'), ['async'] * 3)
        asset = self.client.get(sources[0])
        self.assertEqual(asset.status_code, 200)
        self.assertEqual(b''.join(asset.streaming_content), self.image_bytes)
        self.assertIn('private', asset['Cache-Control'])
        self.assertEqual(len(list(asset_directory(self.sub.file.path).glob('*.png'))), 1)
        self.sub.refresh_from_db(); self.assertIsNone(self.sub.grade)

    def test_missing_asset_is_rebuilt_and_replaced_source_rejects_old_asset_url(self):
        source = self.preview().xpath('//img/@src')[0]
        for image in asset_directory(self.sub.file.path).glob('*.png'):
            image.unlink()
        self.preview()
        asset = self.client.get(source)
        self.assertEqual(asset.status_code, 200); asset.close()
        path = Path(self.sub.file.path)
        path.write_bytes(path.read_bytes() + b'changed-version')
        self.assertEqual(self.client.get(source).status_code, 404)

    @skipUnless(os.name == 'posix', 'Cross-worker file lock is supported on Linux')
    def test_overlapping_prefetch_and_opening_convert_the_document_once(self):
        with patch('pptx.Presentation', wraps=Presentation) as parse:
            with ThreadPoolExecutor(max_workers=2) as pool:
                futures = [pool.submit(convert_office, self.sub.file.path, '.pptx') for _ in range(2)]
                results = [f.result(timeout=10) for f in futures]
        self.assertEqual(parse.call_count, 1)
        self.assertEqual(results[0], results[1])

    def test_assets_and_endpoints_require_the_owner_and_validate_asset_names(self):
        source = self.preview().xpath('//img/@src')[0]
        stranger = User.objects.create_user('fast-review-stranger', password='test-password')
        Teacher.objects.create(user=stranger, full_name='Інший вчитель')
        self.client.force_login(stranger)
        for url in [self.preview_url, self.duplicates_url, source]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 404)
        self.client.logout()
        for url in [self.preview_url, self.duplicates_url, source]:
            self.assertEqual(self.client.get(url).status_code, 302)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse('review_preview_asset', args=[self.sub.pk, 'invalid.png'])).status_code, 404)
        for url in [self.preview_url, self.duplicates_url, source]:
            self.assertEqual(self.client.post(url).status_code, 405)

    def test_selected_attachment_uses_its_own_preview_and_asset_version(self):
        first = SubmissionFile.objects.create(submission=self.sub, original_name='first.txt',
                                               file=SimpleUploadedFile('first.txt', b'First file'))
        second = SubmissionFile.objects.create(submission=self.sub, original_name='second.pptx',
                                               file=SimpleUploadedFile('second.pptx', Path(self.sub.file.path).read_bytes()))
        page = self.preview(file_id=second.pk)
        source = page.xpath('//img/@src')[0]
        self.assertEqual(parse_qs(urlsplit(source).query)['file_id'], [str(second.pk)])
        asset = self.client.get(source)
        self.assertEqual(asset.status_code, 200); asset.close()
        self.assertEqual(self.client.get(source.replace(f'file_id={second.pk}', f'file_id={first.pk}')).status_code, 404)

    def test_failed_conversion_does_not_break_shell_or_modify_grade(self):
        with patch('feed.review_preview.convert_office', side_effect=RuntimeError('converter failed')):
            self.assertEqual(self.client.get(self.preview_url).status_code, 503)
            self.assertEqual(self.client.get(self.view_url).status_code, 200)
        self.sub.refresh_from_db(); self.assertIsNone(self.sub.grade)
        self.assertIn('Відповідь 2', self.preview().text_content())

    def test_duplicate_endpoint_preserves_warning_links_and_collaboration_controls(self):
        AssignmentFile.objects.create(assignment=self.assignment, original_name='Умова.pptx',
            file=SimpleUploadedFile('task.pptx', Path(self.sub.file.path).read_bytes()))
        result = self.client.get(self.duplicates_url).json()
        self.assertIn('Копія матеріалу', html.fromstring(result['chip']).text_content())
        self.assertIn('ignore-plagiarism-checkbox', result['html'])
        self.sub.ignore_plagiarism = True; self.sub.save(update_fields=['ignore_plagiarism'])
        result = self.client.get(self.duplicates_url).json()
        self.assertIn('Збіг дозволено', html.fromstring(result['chip']).text_content())
        self.assertIn('Спільна', html.fromstring(result['group_chip']).text_content())
        self.assertIn('checked', result['html'])

    @override_settings(REVIEW_ASYNC=False)
    def test_global_rollback_restores_original_inline_preview_and_duplicates(self):
        response = self.client.get(self.view_url)
        self.assertContains(response, 'Відповідь 2')
        self.assertContains(response, 'data:image/png;base64,')
        self.assertNotContains(response, 'data-review-preview=')
        self.assertNotContains(response, 'data-review-duplicates=')

    def test_sync_fallback_keeps_active_file_filters_and_old_rendering(self):
        response = self.client.get(self.view_url, {'class': self.group.pk, 'grade_filter': 'ungraded', 'review_sync': '1'})
        self.assertContains(response, 'Відповідь 2')
        self.assertFalse(response.context['preview_pending'])
        self.assertFalse(response.context['duplicates_pending'])

    def test_queue_fetches_ids_without_feedback_and_retains_filtered_neighbors(self):
        other_group = ClassGroup.objects.create(name='8-Б')
        other = Submission.objects.create(assignment=self.assignment, teacher=self.teacher, class_group=self.group,
                                         first_name='Учень', last_name='Другий', ai_feedback='x' * 100000)
        Submission.objects.create(assignment=self.assignment, teacher=self.teacher, class_group=other_group,
                                  first_name='Учень', last_name='Третій')
        with CaptureQueriesContext(connection) as queries:
            response = self.client.get(self.view_url, {'class': self.group.pk})
        self.assertEqual(response.context['queue_total'], 2)
        self.assertEqual(response.context['prev_submission'].pk, other.pk)
        self.assertIsNone(response.context['next_submission'])
        id_queries = [q['sql'] for q in queries if re.match(
            r'^SELECT "feed_submission"\."id"(?: AS "[^"]+")? FROM', q['sql'])]
        self.assertTrue(id_queries)
        self.assertFalse(any('ai_feedback' in q for q in id_queries))
