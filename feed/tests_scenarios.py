"""
Повний набір перевірок 23 обов'язкових сценаріїв з ТЗ (Розділ 11).
"""

import io
import datetime
import json
from django.test import TestCase, Client
from django.contrib.auth.models import User
from django.urls import reverse
from django.utils import timezone
from django.core.files.uploadedfile import SimpleUploadedFile
import docx
from unittest.mock import patch

from feed.models import (
    Teacher, ClassGroup, Subject, Assignment, AssignmentFile,
    Submission, BellSchedule, TeacherLessonSchedule,
    AssignmentScheduleTarget, SubmissionActivityLog
)
from feed.utils import get_teacher_upcoming_notifications, sanitize_html, convert_docx_to_html


class ComprehensiveScenariosTest(TestCase):
    """
    Тестування 23 обов'язкових сценаріїв згідно з вимогами ТЗ.
    """

    def setUp(self):
        # 1. Створюємо викладача з обліковим записом
        self.user = User.objects.create_user(
            username='scenario_teacher',
            password='Password123!',
            is_staff=True,
            is_superuser=True
        )
        self.teacher = Teacher.objects.create(
            user=self.user,
            full_name='Мельник Андрій Васильович'
        )

        # 2. Класи
        self.class_7a = ClassGroup.objects.create(grade=7, letter='А', name='7-А')
        self.class_7b = ClassGroup.objects.create(grade=7, letter='Б', name='7-Б')
        self.class_7v = ClassGroup.objects.create(grade=7, letter='В', name='7-В')
        self.teacher.classes.add(self.class_7a, self.class_7b, self.class_7v)

        # 3. Предмет
        self.subject = Subject.objects.create(name='Інформатика', icon='💻', color='#3b82f6')
        self.teacher.subjects.add(self.subject)

        # 4. Дзвінки (розклад уроків)
        self.bell_1 = BellSchedule.objects.create(lesson_number=1, start_time=datetime.time(9, 0), end_time=datetime.time(9, 45))
        self.bell_2 = BellSchedule.objects.create(lesson_number=2, start_time=datetime.time(10, 0), end_time=datetime.time(10, 45))
        self.bell_3 = BellSchedule.objects.create(lesson_number=3, start_time=datetime.time(11, 0), end_time=datetime.time(11, 45))

        # 5. Тестовий клієнт
        self.client = Client()

    def test_scenario_01_and_02_draft_creation_and_published_status(self):
        """
        Сценарій 1: Створити чернетку завдання.
        Сценарій 2: Переконатися, що воно не вважається опублікованим.
        """
        draft = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Чернетка: Алгоритми',
            description='Опис чернетки',
            status=Assignment.STATUS_DRAFT
        )
        draft.classes.add(self.class_7a)

        # Перевірка статусу
        self.assertNotEqual(draft.status, Assignment.STATUS_PUBLISHED)
        self.assertEqual(draft.status, Assignment.STATUS_DRAFT)
        self.assertIsNone(draft.published_at)

        # Не повинно бути в опублікованих
        published_qs = Assignment.objects.filter(status=Assignment.STATUS_PUBLISHED)
        self.assertNotIn(draft, published_qs)

        # Не відображається у публічній стрічці завдань
        resp = self.client.get(reverse('index'))
        self.assertNotContains(resp, 'Чернетка: Алгоритми')

    def test_scenario_03_and_04_publish_assignment_and_notification_disappears(self):
        """
        Сценарій 3: Опублікувати завдання.
        Сценарій 4: Переконатися, що сповіщення про неопубліковану роботу зникає.
        """
        today = timezone.localtime(timezone.now()).date()
        today_dow = today.weekday() + 1

        TeacherLessonSchedule.objects.create(
            teacher=self.teacher,
            class_group=self.class_7a,
            subject=self.subject,
            day_of_week=today_dow,
            bell_slot=self.bell_2 # 10:00 - 10:45
        )

        # Симулюємо час за 30 хвилин до початку уроку (09:30)
        simulated_now = timezone.make_aware(datetime.datetime.combine(today, datetime.time(9, 30)))

        # До публікації завдання: є сповіщення про відсутнє опубліковане завдання
        notifs_before = get_teacher_upcoming_notifications(self.teacher, now_dt=simulated_now)
        self.assertTrue(any(n['type'] == 'missing_task' and '7-А' in n['message'] for n in notifs_before))

        # Створюємо та публікуємо завдання на цей урок
        assignment = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Опубліковане завдання',
            description='Текст завдання',
            status=Assignment.STATUS_PUBLISHED,
            published_at=simulated_now
        )
        assignment.classes.add(self.class_7a)
        AssignmentScheduleTarget.objects.create(
            assignment=assignment,
            class_group=self.class_7a,
            target_day_of_week=today_dow,
            bell_slot=self.bell_2,
            target_date=today
        )

        # Після публікації: сповіщення про неопубліковану роботу на цей урок зникло!
        notifs_after = get_teacher_upcoming_notifications(self.teacher, now_dt=simulated_now)
        self.assertFalse(any(n['type'] == 'missing_task' and '7-А' in n['message'] for n in notifs_after))

    def test_scenario_05_and_06_sunday_creation_for_monday_lesson_in_calendar(self):
        """
        Сценарій 5: Створити завдання в неділю на урок понеділка.
        Сценарій 6: Переконатися, що в календарі воно відображається в понеділок.
        """
        today = timezone.localtime(timezone.now()).date()
        # Знайдемо найближчу неділю та наступний понеділок
        days_until_sunday = (6 - today.weekday()) % 7
        sunday = today + datetime.timedelta(days=days_until_sunday)
        monday = sunday + datetime.timedelta(days=1)

        # Неділя - дата створення/публікації
        sunday_dt = timezone.make_aware(datetime.datetime.combine(sunday, datetime.time(15, 0)))

        assignment = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Завдання задане в неділю на понеділок',
            description='Підготуватися до практичної',
            status=Assignment.STATUS_PUBLISHED,
            created_at=sunday_dt,
            published_at=sunday_dt
        )
        assignment.classes.add(self.class_7a)
        AssignmentScheduleTarget.objects.create(
            assignment=assignment,
            class_group=self.class_7a,
            target_day_of_week=1, # Понеділок
            bell_slot=self.bell_1,
            target_date=monday
        )

        # Перевірка дати уроку
        self.assertEqual(assignment.get_lesson_date(self.class_7a), monday)

        # Перевірка через календар: фільтр календаря на понеділок повинен показувати завдання
        resp_monday = self.client.get(reverse('index'), {'date': monday.strftime('%Y-%m-%d')})
        self.assertContains(resp_monday, 'Завдання задане в неділю на понеділок')

        # Фільтр календаря на неділю НЕ повинен показувати це завдання
        resp_sunday = self.client.get(reverse('index'), {'date': sunday.strftime('%Y-%m-%d')})
        self.assertNotContains(resp_sunday, 'Завдання задане в неділю на понеділок')

    def test_scenario_07_and_08_multiple_classes_distinct_dates_and_times(self):
        """
        Сценарій 7: Призначити завдання на 2–3 класи з різними уроками.
        Сценарій 8: Переконатися, що кожен клас має свою дату та час.
        """
        monday = datetime.date(2026, 9, 14)
        tuesday = datetime.date(2026, 9, 15)

        assignment = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Python: цикли',
            description='Тема циклів for та while',
            status=Assignment.STATUS_PUBLISHED,
            published_at=timezone.now()
        )
        assignment.classes.add(self.class_7a, self.class_7b, self.class_7v)

        # 7-А: понеділок 10:00 (bell_2)
        AssignmentScheduleTarget.objects.create(
            assignment=assignment,
            class_group=self.class_7a,
            target_day_of_week=1,
            bell_slot=self.bell_2,
            target_date=monday
        )
        # 7-Б: понеділок 11:00 (bell_3)
        AssignmentScheduleTarget.objects.create(
            assignment=assignment,
            class_group=self.class_7b,
            target_day_of_week=1,
            bell_slot=self.bell_3,
            target_date=monday
        )
        # 7-В: вівторок 09:00 (bell_1)
        AssignmentScheduleTarget.objects.create(
            assignment=assignment,
            class_group=self.class_7v,
            target_day_of_week=2,
            bell_slot=self.bell_1,
            target_date=tuesday
        )

        # Перевірка get_all_targets_info()
        targets = assignment.get_all_targets_info()
        self.assertEqual(len(targets), 3)

        info_by_class = {t['class_name']: t for t in targets}
        self.assertIn('7-А', info_by_class)
        self.assertIn('7-Б', info_by_class)
        self.assertIn('7-В', info_by_class)

        self.assertEqual(info_by_class['7-А']['date'], monday)
        self.assertEqual(info_by_class['7-А']['time_str'], '10:00')

        self.assertEqual(info_by_class['7-Б']['date'], monday)
        self.assertEqual(info_by_class['7-Б']['time_str'], '11:00')

        self.assertEqual(info_by_class['7-В']['date'], tuesday)
        self.assertEqual(info_by_class['7-В']['time_str'], '09:00')

        # Перевірка відображення на сторінці завдання
        resp = self.client.get(reverse('assignment_detail', args=[assignment.id]))
        self.assertContains(resp, '7-А')
        self.assertContains(resp, '7-Б')
        self.assertContains(resp, '7-В')

    def test_scenario_09_conducted_lessons_count_only_published(self):
        """
        Сценарій 9: Перевірити підрахунок проведених уроків.
        Враховуються лише опубліковані завдання, чернетки не враховуються.
        """
        past_date = datetime.date(2026, 9, 7) # минулий понеділок

        # 1. Створюємо чернетку на минулий урок
        draft_assignment = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Чернетка на минулий урок',
            status=Assignment.STATUS_DRAFT
        )
        draft_assignment.classes.add(self.class_7a)
        AssignmentScheduleTarget.objects.create(
            assignment=draft_assignment,
            class_group=self.class_7a,
            target_day_of_week=1,
            bell_slot=self.bell_1,
            target_date=past_date
        )

        # Чернетка НЕ повинна давати проведений урок
        self.assertEqual(self.teacher.get_calculated_conducted_lessons(), 0)

        # 2. Створюємо опубліковане завдання на минулий урок
        pub_assignment = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Проведений урок: Вступ',
            status=Assignment.STATUS_PUBLISHED,
            published_at=timezone.make_aware(datetime.datetime.combine(past_date, datetime.time(9, 0)))
        )
        pub_assignment.classes.add(self.class_7b)
        AssignmentScheduleTarget.objects.create(
            assignment=pub_assignment,
            class_group=self.class_7b,
            target_day_of_week=1,
            bell_slot=self.bell_2,
            target_date=past_date
        )

        # Тепер 1 урок проведено
        self.assertEqual(self.teacher.get_calculated_conducted_lessons(), 1)

    def test_scenario_10_11_12_docx_upload_and_student_content_view(self):
        """
        Сценарій 10: Завантажити DOCX із реальним текстом.
        Сценарій 11: Відкрити його від імені учня.
        Сценарій 12: Переконатися, що текст доступний.
        """
        # Створюємо реальний .docx за допомогою python-docx
        doc = docx.Document()
        doc.add_heading('Лабораторна робота з інформатики', level=1)
        doc.add_paragraph('Текст інструкції для учня: створити веб-сторінку за зразком.')
        doc_io = io.BytesIO()
        doc.save(doc_io)
        doc_io.seek(0)

        uploaded_docx = SimpleUploadedFile(
            'lab1_instruction.docx',
            doc_io.read(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

        assignment = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Лабораторна робота DOCX',
            status=Assignment.STATUS_PUBLISHED,
            published_at=timezone.now()
        )
        assignment.classes.add(self.class_7a)

        file_obj = AssignmentFile.objects.create(
            assignment=assignment,
            file=uploaded_docx,
            original_name='lab1_instruction.docx'
        )

        # Перевіряємо роботу парсера docx -> HTML
        html_content, err = convert_docx_to_html(file_obj.file.path)
        self.assertIn('Лабораторна робота з інформатики', html_content)
        self.assertIn('Текст інструкції для учня', html_content)

        # Учень відкриває сторінку завдання
        resp_detail = self.client.get(reverse('assignment_detail', args=[assignment.id]))
        self.assertEqual(resp_detail.status_code, 200)
        self.assertContains(resp_detail, 'lab1_instruction.docx')
        self.assertContains(resp_detail, 'Лабораторна робота з інформатики')

    def test_scenario_13_14_15_16_student_submits_teacher_notification_and_grading(self):
        """
        Сценарій 13: Учень здає роботу.
        Сценарій 14: Перевірити появу сповіщення для вчителя.
        Сценарій 15: Оцінити роботу вручну.
        Сценарій 16: Перевірити зникнення/оновлення сповіщення.
        """
        assignment = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Завдання для перевірки оцінювання',
            status=Assignment.STATUS_PUBLISHED,
            published_at=timezone.now()
        )
        assignment.classes.add(self.class_7a)

        self.assertEqual(self.teacher.pending_reviews_count, 0)

        # 13. Учень здає роботу
        submission = Submission.objects.create(
            assignment=assignment,
            first_name='Оксана',
            last_name='Ковальчук',
            class_group=self.class_7a,
            teacher=self.teacher,
            comment_student='Здаю виконане завдання'
        )

        # 14. Сповіщення для вчителя
        self.assertEqual(self.teacher.pending_reviews_count, 1)
        self.client.login(username='scenario_teacher', password='Password123!')
        resp_poll = self.client.get(reverse('teacher_live_status'))
        data = resp_poll.json()
        self.assertEqual(data['pending_submissions_count'], 1)

        # 15. Вчитель оцінює роботу вручну
        resp_grade = self.client.post(
            reverse('grade_submission', args=[submission.id]),
            {'grade': '11', 'comment': 'Чудово виконано!'}
        )
        self.assertEqual(resp_grade.status_code, 200)

        # 16. Перевірка зникнення сповіщення про неоцінену роботу
        submission.refresh_from_db()
        self.assertEqual(submission.grade, '11')
        self.assertEqual(self.teacher.pending_reviews_count, 0)

        resp_poll2 = self.client.get(reverse('teacher_live_status'))
        data2 = resp_poll2.json()
        self.assertEqual(data2['pending_submissions_count'], 0)

    def test_scenario_17_mass_grading_resolves_all_notifications(self):
        """
        Сценарій 17: Повторити через масове оцінювання.
        Масове оцінювання одночасно прибирає сповіщення для всіх обраних робіт.
        """
        assignment = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Завдання для масового оцінювання',
            status=Assignment.STATUS_PUBLISHED,
            published_at=timezone.now()
        )
        assignment.classes.add(self.class_7a)

        sub1 = Submission.objects.create(
            assignment=assignment,
            first_name='Учень1',
            last_name='Тестовий',
            class_group=self.class_7a,
            teacher=self.teacher
        )
        sub2 = Submission.objects.create(
            assignment=assignment,
            first_name='Учень2',
            last_name='Тестовий',
            class_group=self.class_7a,
            teacher=self.teacher
        )
        sub3 = Submission.objects.create(
            assignment=assignment,
            first_name='Учень3',
            last_name='Тестовий',
            class_group=self.class_7a,
            teacher=self.teacher
        )

        self.assertEqual(self.teacher.pending_reviews_count, 3)

        self.client.login(username='scenario_teacher', password='Password123!')
        resp_mass = self.client.post(
            reverse('mass_grade_submissions'),
            json.dumps({'submission_ids': [sub1.id, sub2.id, sub3.id], 'grade': '12'}),
            content_type='application/json'
        )
        self.assertEqual(resp_mass.status_code, 200)
        res_data = resp_mass.json()
        self.assertTrue(res_data.get('success'))
        self.assertEqual(res_data.get('graded_count'), 3)

        # Перевірка оновлення робіт у БД
        for sub in [sub1, sub2, sub3]:
            sub.refresh_from_db()
            self.assertEqual(sub.grade, '12')

        # Кількість неоцінених стала 0
        self.assertEqual(self.teacher.pending_reviews_count, 0)

    def test_scenario_18_and_19_rich_text_formatting_and_xss_protection(self):
        """
        Сценарій 18: Створити завдання з форматованим текстом.
        Сценарій 19: Перевірити форматування від імені учня (та захист від XSS).
        """
        raw_rich_text = (
            '<h3>Інструкція до роботи</h3>'
            '<p>Виконайте програму на <b>Python</b> з використанням <i>циклів</i> та <u>списків</u>:</p>'
            '<ul><li>Крок 1</li><li>Крок 2</li></ul>'
            '<script>alert("XSS")</script>'
        )

        assignment = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Форматоване завдання',
            description=raw_rich_text,
            status=Assignment.STATUS_PUBLISHED,
            published_at=timezone.now()
        )
        assignment.classes.add(self.class_7a)

        # Перевірка санітизації через get_formatted_description
        formatted = assignment.get_formatted_description()

        # Форматування збережене
        self.assertIn('<b>Python</b>', formatted)
        self.assertIn('<i>циклів</i>', formatted)
        self.assertIn('<u>списків</u>', formatted)
        self.assertIn('<li>Крок 1</li>', formatted)

        # Небезпечний скрипт видалено разом з вмістом
        self.assertNotIn('<script>', formatted)
        self.assertNotIn('alert("XSS")', formatted)

        # Перегляд від імені учня містить безпечне форматування
        resp = self.client.get(reverse('assignment_detail', args=[assignment.id]))
        self.assertContains(resp, '<b>Python</b>')
        self.assertNotContains(resp, 'alert("XSS")')

    def test_scenario_20_21_22_23_assignment_duplication_and_schedule_target(self):
        """
        Сценарій 20: Створити копію завдання.
        Сценарій 21: Переконатися, що створено новий запис, а оригінал не змінений.
        Сценарій 22: Переконатися, що для копії правильно відображається вибраний урок.
        Сценарій 23: Перевірити, що копія потрапила в правильну дату календаря.
        """
        original = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Оригінальне завдання',
            description='<b>Форматований оригінальний текст</b>',
            status=Assignment.STATUS_PUBLISHED,
            published_at=timezone.now()
        )
        original.classes.add(self.class_7a)

        new_lesson_date = datetime.date(2026, 9, 21) # Наступний понеділок

        self.client.login(username='scenario_teacher', password='Password123!')
        resp_dup = self.client.post(
            reverse('assignment_duplicate', args=[original.id]),
            {
                'duplicate_classes': [str(self.class_7b.id)],
                'duplicate_target_date': new_lesson_date.strftime('%Y-%m-%d'),
                'publish_now': '1'
            }
        )
        self.assertEqual(resp_dup.status_code, 302)

        # 21. Створено новий запис, оригінал не змінений
        duplicate = Assignment.objects.exclude(id=original.id).filter(duplicated_from=original).first()
        self.assertIsNotNone(duplicate)
        self.assertNotEqual(duplicate.id, original.id)
        self.assertEqual(duplicate.title, original.title)
        self.assertEqual(duplicate.description, original.description)
        self.assertEqual(list(duplicate.classes.all()), [self.class_7b])
        self.assertEqual(list(original.classes.all()), [self.class_7a])

        # 22. Для копії правильно відображається вибраний урок
        self.assertEqual(duplicate.get_lesson_date(self.class_7b), new_lesson_date)

        # 23. Копія потрапила в правильну дату календаря
        resp_cal_new = self.client.get(reverse('index'), {'date': new_lesson_date.strftime('%Y-%m-%d')})
        self.assertContains(resp_cal_new, 'Оригінальне завдання')

    @patch('django.utils.timezone.now', return_value=timezone.make_aware(datetime.datetime(2026, 9, 22, 10, 0, 0)))
    def test_multi_class_schedule_status_and_card_aging(self, mock_now):
        """
        Тестування вимоги:
        1. Одне завдання для кількох класів з різними уроками.
        2. Відображення статусу кожного класу (актуальний світиться зеленим, пройдений — сірим).
        3. Динамічний перехід актуальності між класами у бейджі.
        4. Блок картки НЕ сіріє, поки не пройшли всі уроки, і стає чорно-білим (card-age-older)
           лише після проходження всіх уроків.
        """
        today = timezone.localtime(timezone.now()).date()
        yesterday = today - datetime.timedelta(days=1)
        tomorrow = today + datetime.timedelta(days=1)
        four_days_ago = today - datetime.timedelta(days=4)

        # Створюємо завдання, опубліковане 5 днів тому
        assignment = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Фізика: Закон Ома для кількох класів',
            description='Лабораторна робота',
            status=Assignment.STATUS_PUBLISHED,
            published_at=timezone.now() - datetime.timedelta(days=5)
        )
        assignment.classes.add(self.class_7a, self.class_7b)

        # 7-А мав урок вчора
        target_7a = AssignmentScheduleTarget.objects.create(
            assignment=assignment,
            class_group=self.class_7a,
            target_date=yesterday,
            target_day_of_week=yesterday.weekday() + 1,
            bell_slot=self.bell_1
        )
        # 7-Б має урок сьогодні (або завтра)
        target_7b = AssignmentScheduleTarget.objects.create(
            assignment=assignment,
            class_group=self.class_7b,
            target_date=tomorrow,
            target_day_of_week=tomorrow.weekday() + 1,
            bell_slot=self.bell_2
        )

        # --- Перевірка старіння картки: хоча публікація була 5 днів тому,
        # уроки ще не закінчились (7-Б на завтра), тому картка свіжа (card-age-today)! ---
        self.assertEqual(assignment.days_since_all_lessons_passed, 0)
        self.assertEqual(assignment.age_card_class, 'card-age-today')

        # --- Перевірка статусів класів у get_all_targets_info() ---
        targets_info = assignment.get_all_targets_info()
        info_by_cls = {t['class_name']: t for t in targets_info}

        # 7-А: пройдений урок -> status 'past', badge_class 'class-status-past' (сірий)
        self.assertEqual(info_by_cls['7-А']['status'], 'past')
        self.assertEqual(info_by_cls['7-А']['badge_class'], 'class-status-past')
        self.assertTrue(info_by_cls['7-А']['is_past'])

        # 7-Б: майбутній урок -> status 'upcoming', badge_class 'class-status-upcoming'
        self.assertEqual(info_by_cls['7-Б']['status'], 'upcoming')
        self.assertEqual(info_by_cls['7-Б']['badge_class'], 'class-status-upcoming')
        self.assertTrue(info_by_cls['7-Б']['is_upcoming'])

        # --- Перевірка динамічного бейджа: для 7-А урок пройшов, але для 7-Б актуальний на завтра ---
        badge = assignment.get_relevance_badge()
        self.assertEqual(badge['badge_class'], 'badge-relevance-upcoming')
        self.assertIn('7-Б', badge['badge_text'])

        # --- Тепер 7-Б має урок СЬОГОДНІ пізніше (наприклад, увечері) ---
        bell_evening = BellSchedule.objects.create(
            lesson_number=8,
            start_time=datetime.time(23, 0),
            end_time=datetime.time(23, 45)
        )
        target_7b.target_date = today
        target_7b.target_day_of_week = today.weekday() + 1
        target_7b.bell_slot = bell_evening
        target_7b.save()

        # Картка залишається сьогоднішньою (0 днів)
        self.assertEqual(assignment.days_since_all_lessons_passed, 0)
        self.assertEqual(assignment.age_card_class, 'card-age-today')

        badge_today = assignment.get_relevance_badge()
        self.assertEqual(badge_today['badge_class'], 'badge-relevance-today')
        self.assertIn('7-Б', badge_today['badge_text'])

        # --- Тепер 7-Б має урок ПРЯМО ЗАРАЗ ---
        now_time = timezone.localtime(timezone.now()).time()
        start_m = max(0, now_time.hour * 60 + now_time.minute - 10)
        end_m = min(23 * 60 + 59, now_time.hour * 60 + now_time.minute + 30)
        bell_now = BellSchedule.objects.create(
            lesson_number=9,
            start_time=datetime.time(start_m // 60, start_m % 60),
            end_time=datetime.time(end_m // 60, end_m % 60)
        )
        target_7b.bell_slot = bell_now
        target_7b.save()

        badge_now = assignment.get_relevance_badge()
        self.assertEqual(badge_now['badge_class'], 'badge-relevance-now')
        self.assertIn('7-Б', badge_now['badge_text'])

        # Статус 7-Б у get_all_targets_info() світиться зеленим (now)
        targets_info_now = assignment.get_all_targets_info()
        info_now = {t['class_name']: t for t in targets_info_now}
        self.assertEqual(info_now['7-Б']['status'], 'now')
        self.assertEqual(info_now['7-Б']['badge_class'], 'class-status-now')
        self.assertTrue(info_now['7-Б']['is_now'])

        # А 7-А залишається сірим (past)
        self.assertEqual(info_now['7-А']['status'], 'past')
        self.assertEqual(info_now['7-А']['badge_class'], 'class-status-past')

        # --- Тепер обидва уроки пройшли 4 дні тому ---
        target_7a.target_date = four_days_ago
        target_7a.save()
        target_7b.target_date = four_days_ago
        target_7b.save()

        # Тепер усі уроки пройшли 4 дні тому -> картка переходить у card-age-older (чорно-біла)
        self.assertEqual(assignment.days_since_all_lessons_passed, 4)
        self.assertEqual(assignment.age_card_class, 'card-age-older')

        # Бейдж відображає завершення всіх уроків
        badge_past = assignment.get_relevance_badge()
        self.assertEqual(badge_past['badge_class'], 'badge-relevance-past')
        self.assertEqual(badge_past['badge_text'], '✓ Всі уроки пройшли')

        # Перевірка рендерингу у головній стрічці
        resp_feed = self.client.get(reverse('index'))
        self.assertEqual(resp_feed.status_code, 200)
        self.assertContains(resp_feed, 'Фізика: Закон Ома для кількох класів')
        self.assertContains(resp_feed, 'class-status-past')

    def test_card_description_markup_rendered_without_raw_tags(self):
        """
        Перевірка, що розмітка (<b>, <strong>, Markdown тощо) коректно рендериться на головній сторінці
        як відформатований текст, а не виводиться як сирі символи тегів (наприклад, <b>лгнаеапнго</b>).
        """
        assignment_html = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Завдання з HTML-розміткою',
            description='<p>Важлива тема: <b>лгнаеапнго</b> та <i>курсив</i></p>',
            status=Assignment.STATUS_PUBLISHED,
            published_at=timezone.now()
        )
        assignment_html.classes.add(self.class_7a)

        assignment_md = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Завдання з Markdown-розміткою',
            description='Тут є **жирний текст** та *курсив*',
            status=Assignment.STATUS_PUBLISHED,
            published_at=timezone.now()
        )
        assignment_md.classes.add(self.class_7a)

        # 1. Метод моделі get_card_description повертає безпечний HTML з тегами
        desc_html = assignment_html.get_card_description()
        self.assertIn('<b>лгнаеапнго</b>', desc_html)
        self.assertIn('<i>курсив</i>', desc_html)
        self.assertNotIn('<p>', desc_html)

        desc_md = assignment_md.get_card_description()
        self.assertIn('<b>жирний текст</b>', desc_md)
        self.assertIn('<i>курсив</i>', desc_md)
        self.assertNotIn('**', desc_md)

        # 2. На головній сторінці теги не ескейпляться як &lt;b&gt; (сирий текст)
        resp = self.client.get(reverse('index'))
        self.assertEqual(resp.status_code, 200)
        content = resp.content.decode('utf-8')
        self.assertNotIn('&lt;b&gt;лгнаеапнго&lt;/b&gt;', content)
        self.assertIn('<b>лгнаеапнго</b>', content)
        self.assertIn('<b>жирний текст</b>', content)

        # 3. У feed_fragment також коректно
        resp_frag = self.client.get(reverse('feed_fragment'))
        self.assertEqual(resp_frag.status_code, 200)
        content_frag = resp_frag.content.decode('utf-8')
        self.assertNotIn('&lt;b&gt;лгнаеапнго&lt;/b&gt;', content_frag)
        self.assertIn('<b>лгнаеапнго</b>', content_frag)

    def test_filtered_submissions_navigation_in_file_viewer(self):
        """
        Тест збереження фільтрів при перевірці робіт у File Viewer:
        - Фільтрація за класом, темою/завданням, датою та статусом
        - Навігація «Попередня» / «Наступна» виключно в межах відфільтрованого списку
        - Відсутність фолбеку на роботи інших класів/завдань при активному фільтрі
        - Повернення «Назад» з точним збереженням активних фільтрів
        """
        # Створюємо спільне завдання для 7-А та 7-Б
        assignment = Assignment.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            title='Практична робота з фільтрацією',
            status=Assignment.STATUS_PUBLISHED,
            published_at=timezone.now()
        )
        assignment.classes.add(self.class_7a, self.class_7b)

        now = timezone.now()
        # 2 здачі для 7-А
        sub_7a_1 = Submission.objects.create(
            assignment=assignment,
            class_group=self.class_7a,
            teacher=self.teacher,
            first_name='Андрій',
            last_name='Шевченко',
            comment_student='Робота 1 7-А',
            submitted_at=now - datetime.timedelta(minutes=30),
            is_latest_attempt=True
        )
        sub_7a_2 = Submission.objects.create(
            assignment=assignment,
            class_group=self.class_7a,
            teacher=self.teacher,
            first_name='Богдан',
            last_name='Хмельницький',
            comment_student='Робота 2 7-А',
            submitted_at=now - datetime.timedelta(minutes=20),
            is_latest_attempt=True
        )
        # 2 здачі для 7-Б
        sub_7b_1 = Submission.objects.create(
            assignment=assignment,
            class_group=self.class_7b,
            teacher=self.teacher,
            first_name='Василь',
            last_name='Стус',
            comment_student='Робота 1 7-Б',
            submitted_at=now - datetime.timedelta(minutes=10),
            is_latest_attempt=True
        )
        sub_7b_2 = Submission.objects.create(
            assignment=assignment,
            class_group=self.class_7b,
            teacher=self.teacher,
            first_name='Григорій',
            last_name='Сковорода',
            comment_student='Робота 2 7-Б',
            submitted_at=now,
            is_latest_attempt=True
        )

        self.client.login(username='scenario_teacher', password='Password123!')

        # 1. Перевірка дашборду всіх здач з фільтром за класом 7-А
        dash_url = f"{reverse('all_submissions_dashboard')}?class={self.class_7a.id}"
        resp_dash = self.client.get(dash_url)
        self.assertEqual(resp_dash.status_code, 200)
        self.assertIn('filter_querystring', resp_dash.context)
        self.assertIn(f'class={self.class_7a.id}', resp_dash.context['filter_querystring'])
        self.assertIn('from=all_submissions', resp_dash.context['filter_querystring'])

        # Посилання «Перевірити» в таблиці містить параметри фільтра
        expected_sub1_url = f"{reverse('view_file', args=[sub_7a_1.id])}?class={self.class_7a.id}&amp;from=all_submissions"
        self.assertContains(resp_dash, expected_sub1_url)

        # 2. Відкриваємо першу роботу (sub_7a_2, оскільки order_by -submitted_at) з фільтром 7-А
        vf_url = f"{reverse('view_file', args=[sub_7a_2.id])}?class={self.class_7a.id}&from=all_submissions"
        resp_vf = self.client.get(vf_url)
        self.assertEqual(resp_vf.status_code, 200)

        # Перевірка контексту черги
        self.assertTrue(resp_vf.context['is_filtered'])
        self.assertEqual(resp_vf.context['queue_total'], 2)
        self.assertEqual(resp_vf.context['queue_pos'], 1)
        self.assertIn('7-А', resp_vf.context['queue_filter_desc'])

        # Наступною має бути sub_7a_1, а не sub_7b_1 чи sub_7b_2!
        self.assertIsNotNone(resp_vf.context['next_submission'])
        self.assertEqual(resp_vf.context['next_submission'].id, sub_7a_1.id)
        self.assertIsNone(resp_vf.context['prev_submission'])

        # Кнопка повернення має вести до all_submissions_dashboard з фільтром 7-А
        self.assertEqual(resp_vf.context['back_url'], f"{reverse('all_submissions_dashboard')}?class={self.class_7a.id}")
        self.assertContains(resp_vf, f"{reverse('all_submissions_dashboard')}?class={self.class_7a.id}")

        # Посилання «Наступна » містить параметри фільтрації
        expected_next_url = f"{reverse('view_file', args=[sub_7a_1.id])}?class={self.class_7a.id}&amp;from=all_submissions"
        self.assertContains(resp_vf, expected_next_url)

        # 3. Переходимо до останньої роботи 7-А (sub_7a_1)
        vf_last_url = f"{reverse('view_file', args=[sub_7a_1.id])}?class={self.class_7a.id}&from=all_submissions"
        resp_vf_last = self.client.get(vf_last_url)
        self.assertEqual(resp_vf_last.status_code, 200)
        self.assertEqual(resp_vf_last.context['queue_pos'], 2)
        self.assertEqual(resp_vf_last.context['prev_submission'].id, sub_7a_2.id)

        # Критично: наступної роботи НЕМАЄ і fallback відключено (не стрибає на 7-Б)!
        self.assertIsNone(resp_vf_last.context['next_submission'])
        self.assertIsNone(resp_vf_last.context['next_submission_fallback'])
        self.assertContains(resp_vf_last, 'Всі роботи у вибраному фільтрі перевірено')

        # 4. Перевірка переходу зі сторінки здач завдання assignment_submissions
        asgn_sub_url = f"{reverse('assignment_submissions', args=[assignment.id])}?class={self.class_7b.id}"
        resp_asgn_sub = self.client.get(asgn_sub_url)
        self.assertEqual(resp_asgn_sub.status_code, 200)
        self.assertIn(f'class={self.class_7b.id}', resp_asgn_sub.context['filter_querystring'])
        self.assertIn(f'assignment={assignment.id}', resp_asgn_sub.context['filter_querystring'])
        self.assertIn('from=assignment_submissions', resp_asgn_sub.context['filter_querystring'])

        # Відкриття роботи 7-Б повертає кнопку назад до assignment_submissions
        vf_7b_url = f"{reverse('view_file', args=[sub_7b_2.id])}?{resp_asgn_sub.context['filter_querystring']}"
        resp_vf_7b = self.client.get(vf_7b_url)
        self.assertEqual(resp_vf_7b.status_code, 200)
        self.assertEqual(
            resp_vf_7b.context['back_url'],
            f"{reverse('assignment_submissions', args=[assignment.id])}?class={self.class_7b.id}&assignment={assignment.id}"
        )
        self.assertEqual(resp_vf_7b.context['queue_total'], 2)
        self.assertEqual(resp_vf_7b.context['next_submission'].id, sub_7b_1.id)

        # 5. Перевірка фільтрації за статусом «ungraded» (без оцінки)
        sub_7a_1.grade = '11'
        sub_7a_1.save(update_fields=['grade'])

        # Тепер серед робіт 7-А без оцінки залишилась лише sub_7a_2
        vf_ungraded_url = f"{reverse('view_file', args=[sub_7a_2.id])}?class={self.class_7a.id}&grade_filter=ungraded"
        resp_ungraded = self.client.get(vf_ungraded_url)
        self.assertEqual(resp_ungraded.status_code, 200)
        self.assertEqual(resp_ungraded.context['queue_total'], 1)
        self.assertIsNone(resp_ungraded.context['next_submission'])
        self.assertIn('Без оцінки', resp_ungraded.context['queue_filter_desc'])

        # 6. Перевірка фільтрації за датою
        today_str = now.strftime('%Y-%m-%d')
        vf_date_url = f"{reverse('view_file', args=[sub_7a_2.id])}?date={today_str}"
        resp_date = self.client.get(vf_date_url)
        self.assertEqual(resp_date.status_code, 200)
        self.assertTrue(resp_date.context['is_filtered'])
        self.assertIn(today_str.split('-')[0], resp_date.context['queue_filter_desc'])

    def test_collective_work_submission_and_grade_sync(self):
        """
        Тест здачі колективної (групової) роботи з кількома співавторами:
        1. Здача роботи учнем з додаванням 2 співавторів через форму.
        2. Перевірка створення зв'язаних робіт у базі даних та позначення як колективна робота.
        3. Перевірка відображення статусу колективної роботи в інтерфейсі перегляду файлів.
        4. Виставлення оцінки вчителем першій роботі -> оцінка автоматично ставиться всім співавторам.
        5. Зміна оцінки через роботу співавтора -> синхронізація оцінки до основної та решти учасників.
        6. Перевірка відображення оцінки для всіх учасників у журналі gradebook.
        """
        from io import BytesIO
        from django.core.files.uploadedfile import SimpleUploadedFile

        assignment = Assignment.objects.create(
            title='Колективний проєкт з інформатики',
            description='Створіть спільний проєкт у групі.',
            teacher=self.teacher,
            subject=self.subject,
            status=Assignment.STATUS_PUBLISHED
        )
        assignment.classes.add(self.class_7a)

        # 1. Учень здає роботу із двома співавторами
        submit_url = reverse('submit_assignment', args=[assignment.id])
        fake_file = SimpleUploadedFile("project.txt", b"Group work contents", content_type="text/plain")

        post_data = {
            'full_name': 'Шевченко Тарас',
            'class_group': self.class_7a.id,
            'files': [fake_file],
            'comment_student': 'Здаємо наш спільний командний проєкт.',
            'coauthors': ['Франко Іван', 'Леся Українка', ''],  # Включаючи порожнє поле для перевірки очистки
        }

        resp = self.client.post(submit_url, post_data)
        self.assertEqual(resp.status_code, 302)

        # 2. Перевіряємо створені роботи
        subs = list(Submission.objects.filter(assignment=assignment, class_group=self.class_7a))
        self.assertEqual(len(subs), 3)

        primary_sub = Submission.objects.get(assignment=assignment, last_name='Шевченко', first_name='Тарас')
        self.assertTrue(primary_sub.is_group_work)
        self.assertTrue(primary_sub.is_collective_work())
        self.assertIsNone(primary_sub.primary_submission)

        coauthors_display = primary_sub.get_group_members_display()
        self.assertIn('Шевченко Тарас', coauthors_display)
        self.assertIn('Франко Іван', coauthors_display)

        # Перевіряємо наявність робіт для співавторів
        franko_sub = Submission.objects.filter(assignment=assignment, last_name='Франко').first()
        self.assertIsNotNone(franko_sub)
        self.assertTrue(franko_sub.is_group_work)
        self.assertTrue(franko_sub.is_collective_work())
        self.assertEqual(franko_sub.primary_submission, primary_sub)

        lesya_sub = Submission.objects.filter(assignment=assignment, last_name__in=['Леся', 'Українка']).first()
        self.assertIsNotNone(lesya_sub)
        self.assertTrue(lesya_sub.is_group_work)
        self.assertEqual(lesya_sub.primary_submission, primary_sub)

        # 3. Вчитель входить у систему та відкриває переглядач роботи первинного автора
        self.client.login(username='scenario_teacher', password='Password123!')
        vf_url = reverse('view_file', args=[primary_sub.id])
        resp_vf = self.client.get(vf_url)
        self.assertEqual(resp_vf.status_code, 200)
        self.assertContains(resp_vf, 'Колективна робота')
        self.assertContains(resp_vf, 'Франко Іван')

        # 4. Вчитель оцінює роботу primary_sub оцінкою «11»
        grade_resp = self.client.post(vf_url, {
            'action': 'grade',
            'grade': '11',
            'submission_id': primary_sub.id,
        }, HTTP_X_REQUESTED_WITH='XMLHttpRequest')

        self.assertEqual(grade_resp.status_code, 200)
        grade_json = grade_resp.json()
        self.assertEqual(grade_json['status'], 'success')
        self.assertEqual(grade_json['grade'], '11')
        self.assertTrue(len(grade_json['coauthors_graded']) >= 2)

        # Перевіряємо оновлення оцінок у базі для всіх трьох учнів
        primary_sub.refresh_from_db()
        franko_sub.refresh_from_db()
        lesya_sub.refresh_from_db()

        self.assertEqual(primary_sub.grade, '11')
        self.assertEqual(franko_sub.grade, '11')
        self.assertEqual(lesya_sub.grade, '11')
        self.assertEqual(franko_sub.graded_by, self.user)
        self.assertEqual(lesya_sub.graded_by, self.user)

        # 5. Зворотна синхронізація: вчитель відкриває роботу Франка і змінює оцінку на «12»
        vf_franko_url = reverse('view_file', args=[franko_sub.id])
        grade_resp2 = self.client.post(vf_franko_url, {
            'action': 'grade',
            'grade': '12',
            'submission_id': franko_sub.id,
        }, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(grade_resp2.status_code, 200)

        primary_sub.refresh_from_db()
        franko_sub.refresh_from_db()
        lesya_sub.refresh_from_db()

        self.assertEqual(primary_sub.grade, '12')
        self.assertEqual(franko_sub.grade, '12')
        self.assertEqual(lesya_sub.grade, '12')

        # 6. Перевіряємо електронний журнал (gradebook)
        gb_url = f"{reverse('gradebook')}?class={self.class_7a.id}"
        resp_gb = self.client.get(gb_url)
        self.assertEqual(resp_gb.status_code, 200)
        self.assertContains(resp_gb, 'Шевченко')
        self.assertContains(resp_gb, 'Франко')
        self.assertContains(resp_gb, '12')





