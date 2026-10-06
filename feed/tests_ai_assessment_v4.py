"""Assessment tests use synthetic work and mocked providers, never real API keys."""
import base64
import io
import json
import os
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from .ai_context import assignment_fingerprint, build_assessment_request, evidence_cache, extract_file_evidence, media_for_provider
from .ai_jobs import enqueue_submission_job
from .forms import AssignmentForm
from .gemini_service import analyze_assignment_task_understanding, evaluate_submission_with_gemini, extract_submission_content
from .middleware import set_has_admin
from .models import AICriteriaPreset, AISettings, Assignment, AssignmentFile, ClassGroup, School, Submission, Teacher


class AssessmentV4Tests(TestCase):
    def setUp(self):
        evidence_cache().clear()
        set_has_admin(True)
        self.user = User.objects.create_superuser('v4_teacher', password='test-password')
        School.objects.create(name='Тестова школа', admin=self.user)
        self.teacher = Teacher.objects.create(user=self.user, full_name='Вчитель')
        self.group = ClassGroup.objects.create(name='5-А', grade=5)
        self.older = ClassGroup.objects.create(name='11-А', grade=11)
        self.teacher.classes.add(self.group, self.older)
        self.preset = AICriteriaPreset.objects.create(name='Результати 5 класу', evaluation_type='nus_gr',
            system_prompt='Оцінюй правильність відомостей.', gr_definitions=json.dumps([
                {'code': 'ГР 1', 'name': 'Зміст'}, {'code': 'ГР 10', 'name': 'Інший результат'}]))
        self.traditional = AICriteriaPreset.objects.create(name='Старша школа', evaluation_type='traditional', system_prompt='Оцінюй зміст і формат.')
        self.assignment = Assignment.objects.create(teacher=self.teacher, title='Бюлетень',
            description='Створи бюлетень про безпечну роботу в інтернеті.', status='published',
            default_ai_preset=self.preset, default_ai_grs=json.dumps(['ГР 1']),
            class_ai_overrides={str(self.older.pk): {'preset_id': self.traditional.pk, 'grs': []}})
        self.assignment.classes.add(self.group, self.older)
        self.sub = Submission.objects.create(assignment=self.assignment, teacher=self.teacher,
            class_group=self.group, first_name='Учень', last_name='Тестовий', grade='11',
            file=SimpleUploadedFile('work.txt', 'Безпечні паролі. Не відкривай підозрілих посилань.'.encode()))
        config = AISettings.get_solo()
        config.api_key = 'synthetic-key'
        config.save()
        self.configs = [{'provider': 'gemini', 'api_key': 'synthetic-key', 'model': 'synthetic-model',
                         'custom_url': '', 'is_backup': False}]

    def evaluate(self, result):
        with patch.object(AISettings, 'get_request_configs', return_value=self.configs), \
             patch('feed.gemini_service.call_ai_api', return_value=(200, json.dumps(result), None, {})) as call:
            response = evaluate_submission_with_gemini(self.sub)
        return response, call

    def result(self, **changes):
        result = {'suggested_grade': '8', 'level': 'Достатній (7-9)', 'summary': 'Зміст правильний.',
                  'feedback_comment': 'Перенеси відомості в бюлетень.', 'grade_explanation': 'Зміст виконано, формат потребує доопрацювання.',
                  'strengths': ['Правила безпеки правильні.'], 'weaknesses': ['Подано презентацію замість бюлетеня.'],
                  'criteria_results': [{'criterion': 'Зміст', 'status': 'completed', 'evidence': 'Правила у файлі.', 'recommendation': ''},
                                       {'criterion': 'Формат бюлетеня', 'status': 'partial', 'evidence': 'Подано слайди.', 'recommendation': 'Створи сторінку бюлетеня.'}],
                  'revision_advice': ['Розмісти текст у колонках бюлетеня.'],
                  'tasks_evaluated': [{'task_num': 1, 'task_title': 'Бюлетень', 'status': 'partial'}], 'gr_results': []}
        result.update(changes)
        return result

    def test_compact_request_preserves_student_work_criteria_and_explicit_custom_prompt(self):
        self.assignment.description='Виконати вправу 6.'
        self.assignment.custom_criteria='Формат — 3 бали; зміст — 9 балів.'
        self.preset.is_system=True
        self.preset.extracted_criteria_text=''
        student='Повна робота учня.\n'+'Код і відповідь. '*1000+'ОСТАННІЙ РЯДОК'
        primary='Матеріал вчителя «task.docx»:\nВправа 1. Незадана.\nВправа 6. Задана повна умова.\nВправа 7. Незадана.'
        refs=''.join(f'📽️ Слайд {i}\nВправа {i}\nУмова на слайді {i}.\n' for i in range(1,9))
        prompt,system=build_assessment_request(self.sub,self.preset,[{'code':'ГР 1','name':'Зміст'}],
                    {},[student],[primary],[refs],[],custom_prompt='Збережи мою особливу інструкцію.',
                    ai_settings=AISettings.get_solo(),compact=True)
        self.assertIn(student,prompt)
        self.assertIn(self.assignment.custom_criteria,prompt)
        self.assertIn('Задана повна умова.',prompt)
        self.assertNotIn('Вправа 7. Незадана.',prompt)
        self.assertNotIn('Умова на слайді 7.',prompt)
        self.assertIn('Збережи мою особливу інструкцію.',system)
        self.assertIn('assessment_blocked=true',system)
        self.assertIn('active_result_groups',system)

    def test_class_policy_is_used_in_student_teacher_and_viewer(self):
        self.sub.class_group = self.older
        self.sub.save()
        from .views import _evaluate_student_submission
        with patch('feed.gemini_service.evaluate_submission_with_gemini') as evaluate:
            _evaluate_student_submission(self.sub)
            self.assertEqual(evaluate.call_args.kwargs['criteria_preset'], self.traditional)
            self.assertEqual(evaluate.call_args.kwargs['selected_gr_codes'], [])
        response, call = self.evaluate(self.result())
        self.assertEqual(response['is_traditional'], True)
        self.assertIn('Старша школа', call.call_args.kwargs['prompt_text'])
        self.client.force_login(self.user)
        page = self.client.get(reverse('view_file', args=[self.sub.pk]))
        self.assertEqual(page.context['default_preset'], self.traditional)
        self.assertTrue(self.sub.is_traditional_grading())

    def test_selected_gr_is_exact_and_does_not_match_gr10(self):
        response, call = self.evaluate(self.result(gr_results=[
            {'code': 'ГР 1', 'name': 'Зміст', 'grade': 8}, {'code': 'ГР 10', 'name': 'Інше', 'grade': 1}]))
        self.assertEqual([gr['code'] for gr in response['gr_results']], ['ГР 1'])
        self.assertEqual(response['suggested_grade'], '8')

    def test_wrong_deliverable_preserves_content_credit_and_improvement_advice(self):
        response, call = self.evaluate(self.result(format_warning='Формат презентації не відповідає завданню про бюлетень.'))
        self.assertEqual(response['suggested_grade'], '8')
        self.assertEqual(response['tasks_completed_count'], 0)
        self.assertIn('презентація замість бюлетеня', call.call_args.kwargs['system_prompt'].lower())
        self.sub.refresh_from_db()
        self.assertIn('Чому така оцінка', self.sub.ai_feedback)
        self.assertIn('Створи сторінку бюлетеня', self.sub.ai_feedback)
        self.assertIn('Як покращити', self.sub.ai_feedback)
        self.assertEqual(self.sub.grade, '11')

    def test_model_cannot_manufacture_completed_criterion_from_a_high_grade(self):
        self.assignment.custom_criteria = 'Окремий критерій: зазначити два джерела.'
        self.assignment.save()
        response, _ = self.evaluate(self.result(suggested_grade='12', criteria_results=[]))
        row = next(r for r in response['criteria_results'] if 'два джерела' in r['criterion'])
        self.assertEqual(row['status'], 'unverifiable')
        self.assertNotIn('Підтверджено', row['evidence'])

    def test_blocked_and_malformed_evaluations_do_not_use_student_attempt(self):
        for result in [self.result(assessment_blocked=True), {'feedback_comment': 'Без оцінки'}]:
            with patch('feed.views._evaluate_student_submission', return_value={'status': 'failed'}):
                job = enqueue_submission_job(self.sub, 'student_check')
            self.assertEqual(job.status, 'failed')
            self.assertFalse(job.reservations.filter(used=True).exists())
            response, _ = self.evaluate(result)
            self.assertEqual(response['status'], 'failed')
        self.sub.refresh_from_db()
        self.assertFalse(self.sub.student_ai_checked)
        self.assertEqual(self.sub.grade, '11')
        with patch.object(AISettings, 'get_request_configs', return_value=self.configs), \
             patch('feed.gemini_service.call_ai_api', return_value=(200, 'suggested_grade: 12 broken', None, {})):
            self.assertEqual(evaluate_submission_with_gemini(self.sub)['status'], 'failed')

    def test_successful_self_check_is_concise_and_cannot_repeat_after_resubmission(self):
        from .views import _perform_student_ai_check
        with patch('feed.views._evaluate_student_submission', return_value=dict(self.result(), status='success')):
            job = enqueue_submission_job(self.sub, 'student_check')
        self.assertEqual(job.status, 'succeeded')
        self.assertNotIn('Чому така оцінка', job.result['feedback'])
        self.sub.refresh_from_db()
        # Mocked results still persist the useful explanations through the view.
        self.assertTrue(self.sub.student_ai_checked)
        self.assertNotIn('Чому така оцінка', self.sub.student_ai_feedback)
        self.assertEqual(self.sub.grade, '11')
        session = self.client.session
        session['last_submission_id'] = self.sub.pk
        session.save()
        page = self.client.get(reverse('submit_success', args=[self.assignment.pk]))
        self.assertContains(page, 'Чому така оцінка')
        self.assertContains(page, 'Похвала за роботу')
        self.assertContains(page, 'Рекомендація учню')
        self.assertContains(page, 'Як покращити роботу')
        self.assertNotContains(page, '📋 Перевірка критеріїв')
        self.assertContains(page, '💡')
        self.assertContains(page, 'Розмісти текст у колонках бюлетеня')
        self.assertEqual(self.sub.get_student_ai_evidence_sections(), [])
        later = Submission.objects.create(assignment=self.assignment, teacher=self.teacher, class_group=self.group,
                                          first_name=self.sub.first_name, last_name=self.sub.last_name)
        from .ai_jobs import SelfCheckUnavailable
        with self.assertRaises(SelfCheckUnavailable):
            enqueue_submission_job(later, 'student_check')

    def test_form_saves_only_assigned_classes_and_valid_grs(self):
        data = {'title': 'Робота', 'description': 'Створи документ.', 'classes': [self.group.pk],
                'publish_choice': 'draft', f'class_ai_preset_{self.group.pk}': self.preset.pk,
                f'class_ai_grs_{self.group.pk}': ['ГР 1'], f'class_ai_preset_{self.older.pk}': self.traditional.pk}
        form = AssignmentForm(teacher=self.teacher, data=data)
        self.assertTrue(form.is_valid(), form.errors)
        assignment = form.save_with_status(self.teacher)
        self.assertEqual(assignment.class_ai_overrides, {str(self.group.pk): {'preset_id': self.preset.pk, 'grs': ['ГР 1']}})
        data[f'class_ai_grs_{self.group.pk}'] = ['ГР 100']
        self.assertFalse(AssignmentForm(teacher=self.teacher, data=data).is_valid())

    def test_late_exercise_is_not_truncated_and_parse_is_cached(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp, 'material.txt')
            path.write_text('Теорія\n' * 5000 + 'Вправа 3. Створи бюлетень з двома колонками.', encoding='utf-8')
            first = extract_file_evidence(str(path), visual=False)
            self.assertIn('Вправа 3', first['text'])
            with patch('feed.ai_context._full_text', side_effect=AssertionError('must use cache')):
                self.assertEqual(extract_file_evidence(str(path), visual=False), first)
            path.write_text('Вправа 3. Створи таблицю.', encoding='utf-8')
            self.assertIn('Створи таблицю', extract_file_evidence(str(path), visual=False)['text'])

    def test_full_excel_formulas_charts_and_late_sheet_are_available(self):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp, 'work.xlsx')
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(['Категорія', 'Кількість'])
            sheet.append(['А', 4]); sheet.append(['Б', 7]); sheet['B4'] = '=SUM(B2:B3)'
            chart = BarChart(); chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
            chart.title = 'Порівняння'; sheet.add_chart(chart, 'D1')
            for i in range(6): workbook.create_sheet(f'Аркуш{i}')
            workbook.worksheets[-1]['A1'] = 'Останній аркуш: висновок'
            workbook.save(path)
            evidence = extract_file_evidence(str(path), visual=False)
            self.assertIn('=SUM(B2:B3)', evidence['text'])
            self.assertIn('Порівняння', evidence['text'])
            self.assertIn('Останній аркуш', evidence['text'])

    def test_scratch_keeps_all_blocks_and_reports_static_analysis(self):
        project = {'targets': [{'name': 'Кіт', 'blocks': {'b': {'opcode': 'control_repeat', 'inputs': {'TIMES': [1, [4, '10']]}}},
                                'variables': {'v': ['Рахунок', 1]}, 'lists': {'l': ['Список', [1, 2]]}}]}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp, 'work.sb3')
            with zipfile.ZipFile(path, 'w') as archive:
                archive.writestr('project.json', json.dumps(project))
            data = extract_file_evidence(str(path))
            self.assertIn('control_repeat', data['text'])
            self.assertIn('Рахунок', data['text'])
            self.assertIn('статично', ' '.join(data['limitations']))

    def test_pdf_is_sent_as_images_to_other_providers_and_no_page_is_skipped(self):
        from pypdf import PdfWriter
        buffer = io.BytesIO(); writer = PdfWriter()
        writer.add_blank_page(width=200, height=200); writer.add_blank_page(width=200, height=200); writer.write(buffer)
        media = [{'mime_type': 'application/pdf', 'data': base64.b64encode(buffer.getvalue()).decode(), 'source': 'Умова'}]
        pages = media_for_provider(media, 'openai')
        self.assertEqual(len(pages), 2)
        self.assertTrue(all(p['mime_type'] == 'image/jpeg' for p in pages))
        self.assertIn('сторінка 2', pages[1]['source'])
        self.assertEqual(media_for_provider(media, 'gemini'), media)

    def test_understanding_uses_visual_material_and_cache_invalidates_with_rubric(self):
        file = AssignmentFile.objects.create(assignment=self.assignment, file=SimpleUploadedFile('instructions.txt', 'Вправа 3: створити бюлетень.'.encode()))
        self.assignment.description = 'Виконати вправу 3.'; self.assignment.save()
        data = {'tasks_total_count': 1, 'tasks': [{'num': 3, 'title': 'Бюлетень', 'source': 'instructions.txt', 'expected_actions': 'Створи бюлетень.'}],
                'student_explanation': 'Крок 1. Відкрий матеріал. Крок 2. Створи бюлетень із двома колонками. Крок 3. Перевір і надішли файл.',
                'deliverable': {'format': 'Бюлетень'}, 'teacher_recommendations': []}
        with patch.object(AISettings, 'get_request_configs', return_value=self.configs), \
             patch('feed.gemini_service.call_ai_api', return_value=(200, json.dumps(data), None, {})) as call:
            result = analyze_assignment_task_understanding(self.assignment)
            self.assertIn('instructions.txt', call.call_args.kwargs['prompt_text'])
            self.assertEqual(result['data']['deliverable']['format'], 'Бюлетень')
            analyze_assignment_task_understanding(self.assignment)
            self.assertEqual(call.call_count, 1)
            self.assignment.custom_criteria = 'Дві колонки — 3 бали.'; self.assignment.save()
            analyze_assignment_task_understanding(self.assignment)
            self.assertEqual(call.call_count, 2)

    def test_assignment_page_does_not_parse_all_materials_before_opening(self):
        AssignmentFile.objects.create(assignment=self.assignment, file=SimpleUploadedFile('big.docx', b'synthetic'))
        with patch('feed.views.convert_docx_to_html', side_effect=AssertionError('must be on demand')):
            self.assertEqual(self.client.get(reverse('assignment_detail', args=[self.assignment.pk])).status_code, 200)

    def test_unreadable_only_work_has_no_automatic_grade(self):
        self.sub.file.save('work.mp4', SimpleUploadedFile('work.mp4', b'\x00\x00binary\xff'))
        text, media, error = extract_submission_content(self.sub)
        self.assertEqual((text, media), ([], []))
        self.assertIn('ручна перевірка', error)

    def test_zip_policy_roundtrip_keeps_classes_rubrics_and_source_flag(self):
        from .assignment_ai_policy import export_ai_policy, import_ai_policy
        document = SimpleUploadedFile('criteria.txt', 'Формат: 3 бали. Зміст: 9 балів.'.encode())
        self.traditional.document_file.save('criteria.txt', document)
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            policy = export_ai_policy(self.assignment, archive)
        clone = Assignment.objects.create(teacher=self.teacher, title='Копія', description='Тест')
        clone.classes.add(self.group, self.older)
        with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as archive:
            import_ai_policy(clone, policy, archive)
        self.assertEqual(clone.get_ai_policy(self.older)[0].name, self.traditional.name)
        self.assertEqual(clone.get_ai_policy(self.group)[1], ['ГР 1'])
        self.assertTrue(clone.get_ai_policy(self.older)[0].document_file)

    def test_custom_points_are_not_replaced_by_result_group_average(self):
        self.assignment.custom_criteria = 'Зміст — 9 балів; формат — 3 бали.'
        self.assignment.save()
        response, _ = self.evaluate(self.result(suggested_grade='9', gr_results=[{'code': 'ГР 1', 'grade': 8}]))
        self.assertEqual(response['suggested_grade'], '9')

    def test_compact_context_preserves_group_definition_without_repeated_interpretations(self):
        response, call = self.evaluate(self.result())
        prompt = call.call_args.kwargs['prompt_text']
        self.assertIn('result_group_definitions', prompt)
        self.assertIn('Зміст', prompt)
        self.assertNotIn('final_task_understanding', prompt)
        self.assertNotIn('forbidden_assumptions', prompt)

    def test_import_does_not_reuse_a_different_visual_rubric(self):
        from .assignment_ai_policy import export_ai_policy, import_ai_policy
        self.traditional.document_file.save('rubric.txt', SimpleUploadedFile('rubric.txt', b'original criteria'))
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            policy = export_ai_policy(self.assignment, archive)
        self.traditional.document_file.save('rubric.txt', SimpleUploadedFile('rubric.txt', b'different criteria'))
        clone = Assignment.objects.create(teacher=self.teacher, title='Копія', description='Тест')
        clone.classes.add(self.older)
        with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as archive:
            import_ai_policy(clone, policy, archive)
        imported = clone.get_ai_policy(self.older)[0]
        self.assertNotEqual(imported.pk, self.traditional.pk)
        with imported.document_file.open('rb') as document:
            self.assertEqual(document.read(), b'original criteria')

    def test_slide_number_is_not_an_exercise_number(self):
        from .gemini_service import find_question_by_task_num
        self.assertIsNone(find_question_by_task_num(['📽️ Слайд 3/3: «Вправа 4» Створи таблицю.'], 3))
        self.assertEqual(find_question_by_task_num(['[Слайд 8] 3. Створи бюлетень.'], 3), '[Слайд 8] 3. Створи бюлетень.')

    def test_teacher_link_content_is_cached_and_link_changes_invalidate_task_guide(self):
        from .ai_context import teacher_materials
        from .models import AssignmentLink
        link = AssignmentLink.objects.create(assignment=self.assignment, url='https://example.com/exercise', label='Умова')
        before = assignment_fingerprint(self.assignment)
        with patch('feed.gemini_service.fetch_url_content', return_value=('Умова', 'Вправа 3: створи бюлетень.', None)) as fetch:
            _, reference, _, coverage = teacher_materials(self.assignment)
            teacher_materials(self.assignment)
            self.assertEqual(fetch.call_count, 1)
        self.assertIn('Вправа 3', '\n'.join(reference))
        self.assertTrue(coverage[-1]['limitations'])
        link.label = 'Нова умова'; link.save()
        self.assertNotEqual(before, assignment_fingerprint(self.assignment))

    def test_binary_url_is_not_misread_as_text(self):
        from .gemini_service import fetch_url_content
        response = type('Response', (), {'headers': {'Content-Type': 'application/pdf'}, 'read': lambda self, count: b'%PDF binary',
                                        '__enter__': lambda self: self, '__exit__': lambda *args: None})()
        with patch('feed.safe_http.validate_public_url'), patch('feed.safe_http.public_urlopen', return_value=response):
            _, content, error = fetch_url_content('https://example.com/lesson.pdf')
        self.assertIsNone(content)
        self.assertIn('прикріпіть сам файл', error)

    def test_visual_task_guide_does_not_inherit_wrong_local_deliverable_or_generic_points(self):
        self.assignment.description = 'Виконай вправу 3 зі прикріпленої презентації.'
        self.assignment.custom_criteria = 'Правильний зміст — 9 балів; формат бюлетеня — 3 бали.'
        self.assignment.save()
        answer = {'tasks_total_count': 1, 'student_explanation': 'Крок 1: створи бюлетень. Крок 2: додай правила. Крок 3: збережи документ.',
                  'task_type': 'bulletin', 'deliverable': {'type': 'bulletin', 'description': 'Бюлетень', 'format': 'Документ у дві колонки'},
                  'submission_format_expected': 'Документ у дві колонки',
                  'teacher_requirements': ['Три пояснення правил'],
                  'tasks': [{'num': 1, 'title': 'Бюлетень', 'expected_actions': 'Додай три правила.', 'source': 'Слайд 2'}],
                  'grading_breakdown': {'partial_one_task': '4-5 за відсутність колонок', 'rules': []}}
        with patch.object(AISettings, 'get_request_configs', return_value=self.configs), patch('feed.gemini_service.call_ai_api', return_value=(200, json.dumps(answer), None, {})):
            response = analyze_assignment_task_understanding(self.assignment, force_refresh=True)['data']
        self.assertEqual(response['tasks'][0]['num'], 3)
        self.assertEqual(response['task_interpretation']['task_type'], 'bulletin')
        self.assertEqual(response['final_task_understanding']['expected_result'], 'Документ у дві колонки')
        self.assertEqual(response['teacher_requirements'], ['Три пояснення правил'])
        self.assertNotIn('4-5', response['grading_breakdown']['partial_one_task'])
        self.assertIn(self.assignment.custom_criteria, response['grading_breakdown']['rules'])

    def test_late_presentation_slide_and_chart_are_not_skipped(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        presentation = Presentation()
        for index in range(45):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
            box.text = f'Слайд {index + 1}'
        box.text = 'Вправа 3. Створити бюлетень з діаграмою.'
        chart_data = CategoryChartData(); chart_data.categories = ['А', 'Б']; chart_data.add_series('Показник', [4, 9])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(5), Inches(3), chart_data)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp, 'lesson.pptx'); presentation.save(path)
            evidence = extract_file_evidence(str(path), visual=False)
        self.assertIn('Вправа 3', evidence['text'])
        self.assertIn('Показник', evidence['text'])
        self.assertIn('45/45', evidence['text'])
