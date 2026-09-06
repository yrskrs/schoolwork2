"""
Утиліти для конвертації та обробки файлів і посилань для попереднього перегляду в SchoolNet+
"""
import os
import re
import html
import base64
import subprocess
import shutil
import urllib.request
from typing import Optional, Tuple
from urllib.parse import urlparse, parse_qs
from django.utils.encoding import iri_to_uri

_URL_TITLE_CACHE = {}


def fetch_url_title(url: str, timeout: float = 2.0) -> str:
    """
    Отримує заголовок вебсторінки (<title>) за вказаним URL.
    Використовує кеш у пам'яті для миттєвого доступу без повторних запитів.
    """
    if not url or not isinstance(url, str):
        return ""
    url = url.strip()
    if not (url.startswith('http://') or url.startswith('https://')):
        return ""
    if url in _URL_TITLE_CACHE:
        return _URL_TITLE_CACHE[url]

    try:
        clean_url = iri_to_uri(url)
        req = urllib.request.Request(
            clean_url,
            headers={
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36 SchoolNet/2.0'
            }
        )
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            content_type = resp.headers.get('content-type', '')
            if 'text/html' not in content_type.lower():
                return ""
            raw_bytes = resp.read(40960)
            charset = resp.headers.get_content_charset() or 'utf-8'
            raw_html = raw_bytes.decode(charset, errors='replace')
            m = re.search(r'<title[^>]*>(.*?)</title>', raw_html, re.IGNORECASE | re.DOTALL)
            if m:
                extracted = html.unescape(m.group(1)).strip()
                clean_title = re.sub(r'\s+', ' ', extracted)[:180]
                if clean_title:
                    _URL_TITLE_CACHE[url] = clean_title
                    return clean_title
    except Exception:
        pass

    try:
        parsed = urlparse(url)
        domain = (parsed.netloc or parsed.path).replace('www.', '')
        if domain:
            _URL_TITLE_CACHE[url] = domain
            return domain
    except Exception:
        pass

    return ""


def convert_docx_to_html(file_path: str) -> Tuple[str, Optional[str]]:
    """
    Конвертує .docx або .doc файл у HTML для відображення в браузері.
    Підтримує сучасні .docx (через mammoth / python-docx) та класичні .doc (через antiword / OLE-парсинг).
    """
    if not file_path or not os.path.exists(file_path):
        return "", "Файл не знайдено на сервері"

    is_doc = False
    try:
        with open(file_path, 'rb') as f_check:
            magic = f_check.read(8)
            if magic.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1') or file_path.lower().endswith('.doc'):
                is_doc = True
    except Exception:
        pass

    # 1. ОБРОБКА КЛАСИЧНИХ .doc ФАЙЛІВ (Word 97-2003)
    if is_doc:
        # Спроба 1: Конвертація .doc -> .docx через LibreOffice для 100% збереження структури та таблиць
        converted_docx = None
        try:
            if shutil.which('libreoffice') or shutil.which('soffice'):
                from django.conf import settings
                cache_dir = os.path.join(settings.MEDIA_ROOT, 'previews')
                os.makedirs(cache_dir, exist_ok=True)
                import hashlib
                h_hash = hashlib.md5(file_path.encode('utf-8')).hexdigest()
                cached_docx = os.path.join(cache_dir, f"doc_{h_hash}.docx")
                
                if os.path.exists(cached_docx) and os.path.getmtime(cached_docx) >= os.path.getmtime(file_path):
                    converted_docx = cached_docx
                else:
                    cmd = ['libreoffice', '--headless', '--convert-to', 'docx', '--outdir', cache_dir, file_path]
                    proc = subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=25)
                    raw_docx = os.path.join(cache_dir, os.path.splitext(os.path.basename(file_path))[0] + '.docx')
                    if os.path.exists(raw_docx):
                        if raw_docx != cached_docx:
                            if os.path.exists(cached_docx):
                                os.remove(cached_docx)
                            os.rename(raw_docx, cached_docx)
                        converted_docx = cached_docx
        except Exception:
            converted_docx = None

        if converted_docx and os.path.exists(converted_docx):
            file_path = converted_docx
            is_doc = False  # Переходимо до обробки як сучасного .docx через mammoth нижче!
        else:
            # Спроба 2: antiword з підтримкою DocBook XML (-x db), що зберігає цілісність таблиць
            if shutil.which('antiword'):
                try:
                    proc = subprocess.run(
                        ['antiword', '-x', 'db', file_path],
                        capture_output=True,
                        text=True,
                        errors='replace',
                        timeout=15
                    )
                    if proc.returncode == 0 and proc.stdout.strip():
                        from lxml import etree
                        xml_clean = re.sub(r'<!DOCTYPE[^>]*>', '', proc.stdout, flags=re.DOTALL)
                        xml_clean = re.sub(r'^XML doesn[^\n]*\n', '', xml_clean)
                        parser = etree.XMLParser(recover=True, encoding='utf-8')
                        root = etree.fromstring(xml_clean.encode('utf-8'), parser=parser)

                        def _elem_to_html(elem):
                            text = elem.text or ''
                            res = html.escape(text)
                            for ch in elem:
                                if ch.tag == 'emphasis':
                                    role = ch.get('role', '')
                                    t = 'strong' if role == 'bold' else 'em'
                                    res += f'<{t}>{_elem_to_html(ch)}</{t}>'
                                else:
                                    res += _elem_to_html(ch)
                                if ch.tail:
                                    res += html.escape(ch.tail)
                            return res

                        html_out = ['<div class="document-page">']
                        for child in root.iter():
                            if child.tag == 'para' and child.getparent().tag == 'chapter':
                                c = _elem_to_html(child).strip()
                                if c:
                                    html_out.append(f'<p style="margin-bottom:8px;line-height:1.6;">{c}</p>')
                            elif child.tag == 'informaltable':
                                tgroup = child.find('tgroup')
                                total_cols = int(tgroup.get('cols', '1')) if tgroup is not None else 1
                                html_out.append('<div class="docx-table-wrapper"><table class="docx-table">')
                                for i, row in enumerate(child.iter('row')):
                                    tag = 'th' if i == 0 else 'td'
                                    entries = row.findall('entry')
                                    html_out.append('<tr>')
                                    for e in entries:
                                        paras = [_elem_to_html(p).strip() for p in e.findall('para') if _elem_to_html(p).strip()]
                                        cell_content = '<br>'.join(paras) if paras else _elem_to_html(e).strip()
                                        colspan = f' colspan="{total_cols}"' if len(entries) == 1 and total_cols > 1 else ''
                                        html_out.append(f'<{tag}{colspan}>{cell_content}</{tag}>')
                                    html_out.append('</tr>')
                                html_out.append('</table></div>')
                        html_out.append('</div>')
                        return '\n'.join(html_out), None
                except Exception:
                    pass

            # Спроба 3: Резервний варіант для .doc (видобування тексту з бінарного OLE)
            try:
                with open(file_path, 'rb') as f:
                    content = f.read()
                utf16_runs = re.findall(rb'(?:[\x20-\x7e\xa0-\xff\x00-\x04]\x00|\r\x00|\n\x00){4,}', content)
                lines = []
                for r in utf16_runs:
                    try:
                        t = r.decode('utf-16le', errors='ignore').strip()
                        if len(t) > 3 and not t.startswith(('Normal', 'Heading', 'Title', 'Default')):
                            lines.append(html.escape(t))
                    except Exception:
                        pass
                if lines:
                    html_out = ['<div class="document-page">']
                    for l in lines:
                        html_out.append(f'<p style="margin-bottom:8px;line-height:1.7;">{l}</p>')
                    html_out.append('</div>')
                    return '\n'.join(html_out), None
            except Exception as e:
                return "", f"Помилка відкриття .doc документа: {str(e)}"

    # 2. ОБРОБКА СУЧАСНИХ .docx ФАЙЛІВ
    # Спроба 1: mammoth (найкраща якість конвертації, зі зображеннями як base64)
    try:
        import mammoth

        def _convert_image_to_base64(image):
            """Конвертує вбудоване зображення Word документу в base64 data URI."""
            with image.open() as img_bytes:
                encoded = base64.b64encode(img_bytes.read()).decode('ascii')
            ct = image.content_type or 'image/png'
            return {'src': f'data:{ct};base64,{encoded}'}

        style_map = """
p[style-name='Heading 1'] => h1.doc-h1:fresh
p[style-name='Heading 2'] => h2.doc-h2:fresh
p[style-name='Heading 3'] => h3.doc-h3:fresh
p[style-name='Заголовок 1'] => h1.doc-h1:fresh
p[style-name='Заголовок 2'] => h2.doc-h2:fresh
p[style-name='Заголовок 3'] => h3.doc-h3:fresh
table => table.docx-table:fresh
r[style-name='Strong'] => strong
b => strong
i => em
u => u
strike => s
"""

        with open(file_path, 'rb') as docx_file:
            result = mammoth.convert_to_html(
                docx_file,
                convert_image=mammoth.images.img_element(_convert_image_to_base64),
                style_map=style_map,
            )
            if result.value and result.value.strip():
                content = result.value
                if '<table' in content:
                    content = re.sub(r'(<table\b[^>]*>.*?</table>)', r'<div class="docx-table-wrapper">\1</div>', content, flags=re.DOTALL)
                return f'<div class="document-page">{content}</div>', None
    except Exception:
        pass

    # Спроба 2: python-docx (якщо mammoth недоступний)
    # Обробляємо елементи у порядку документу (параграфи + таблиці разом)
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from lxml import etree

        doc = Document(file_path)
        html_parts = ['<div class="document-page">']

        def _render_para(para):
            if not para.text.strip():
                return '<p>&nbsp;</p>'
            style_name = getattr(getattr(para, 'style', None), 'name', '') or ''
            text_escaped = html.escape(para.text)
            sl = style_name.lower()
            if 'heading 1' in sl or 'заголовок 1' in sl:
                return f'<h1 class="doc-h1">{text_escaped}</h1>'
            elif 'heading 2' in sl or 'заголовок 2' in sl:
                return f'<h2 class="doc-h2">{text_escaped}</h2>'
            elif 'heading' in sl or 'заголовок' in sl:
                return f'<h3 class="doc-h3">{text_escaped}</h3>'
            else:
                return f'<p style="margin-bottom:8px;line-height:1.7;">{text_escaped}</p>'

        def _render_table(table):
            rows_html = []
            for i, row in enumerate(table.rows):
                tag = 'th' if i == 0 else 'td'
                cells_html = ''.join(
                    f'<{tag}>{html.escape(cell.text.strip())}</{tag}>'
                    for cell in row.cells
                )
                rows_html.append(f'<tr>{cells_html}</tr>')
            return '<div class="docx-table-wrapper"><table class="docx-table">' + ''.join(rows_html) + '</table></div>'

        # Iterate in document order via the body's XML children
        from docx.text.paragraph import Paragraph
        from docx.table import Table as DocxTable

        body = doc.element.body
        for child in body:
            tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
            if tag == 'p':
                para = Paragraph(child, body)
                html_parts.append(_render_para(para))
            elif tag == 'tbl':
                table = DocxTable(child, body)
                html_parts.append(_render_table(table))

        html_parts.append('</div>')
        html_content = '\n'.join(html_parts)
        if html_content.strip() == '<div class="document-page"></div>':
            html_content = '<div class="document-page"><p class="text-muted">Документ порожній або не містить текстового контенту.</p></div>'
        return html_content, None
    except Exception as e:
        error_msg = f"Помилка при читанні документа: {str(e)}"
        return "", error_msg



def convert_xlsx_to_html(file_path: str, max_rows: int = 100) -> Tuple[str, Optional[str]]:
    """
    Конвертує .xlsx файл у HTML таблиці з підтримкою формул та стилів.
    """
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill
        from openpyxl.utils import get_column_letter
        
        wb_values = load_workbook(file_path, data_only=True)
        wb_formulas = load_workbook(file_path, data_only=False)
        
        html_parts = []
        
        for sheet_name in wb_values.sheetnames:
            sheet_values = wb_values[sheet_name]
            sheet_formulas = wb_formulas[sheet_name]
            
            html_parts.append('<div class="document-page excel-sheet mb-4">')
            html_parts.append(f'<h4 class="sheet-title mb-3" style="color:var(--color-primary, #6366f1);font-weight:700;">📊 Аркуш: {sheet_name}</h4>')
            
            max_row = min(sheet_values.max_row or 1, max_rows)
            max_col = sheet_values.max_column or 1
            
            if max_row == 0 or max_col == 0:
                html_parts.append('<p class="text-muted">Аркуш порожній</p>')
                html_parts.append('</div>')
                continue
            
            html_parts.append('<div class="table-responsive" style="overflow-x:auto;">')
            html_parts.append('<table class="excel-table table table-bordered" style="width:100%;border-collapse:collapse;font-size:13px;">')
            
            # Column headers (A, B, C...)
            html_parts.append('<thead><tr style="background:var(--color-bg-secondary,#f8fafc);"><th class="row-header" style="width:40px;text-align:center;">#</th>')
            for col in range(1, max_col + 1):
                col_letter = get_column_letter(col)
                html_parts.append(f'<th style="text-align:center;padding:6px 10px;font-weight:600;">{col_letter}</th>')
            html_parts.append('</tr></thead>')
            
            html_parts.append('<tbody>')
            
            for row_idx in range(1, max_row + 1):
                html_parts.append('<tr>')
                html_parts.append(f'<td class="row-header" style="background:var(--color-bg-secondary,#f8fafc);text-align:center;font-weight:600;width:40px;">{row_idx}</td>')
                
                for col_idx in range(1, max_col + 1):
                    cell_val = sheet_values.cell(row=row_idx, column=col_idx)
                    cell_formula = sheet_formulas.cell(row=row_idx, column=col_idx)
                    
                    value = cell_val.value
                    if value is None:
                        value = ''
                    
                    formula = ""
                    if cell_formula.data_type == 'f':
                        formula = f"={cell_formula.value}"
                    elif value != "":
                        formula = str(value)
                        
                    style_parts = ['padding:6px 10px;', 'border:1px solid var(--color-border, #e2e8f0);']
                    
                    if cell_val.font:
                        if cell_val.font.bold:
                            style_parts.append('font-weight: bold;')
                        if cell_val.font.italic:
                            style_parts.append('font-style: italic;')
                        if cell_val.font.color and hasattr(cell_val.font.color, 'rgb') and cell_val.font.color.rgb:
                            color = str(cell_val.font.color.rgb)
                            if len(color) == 8:
                                color = '#' + color[2:]
                            elif len(color) == 6:
                                color = '#' + color
                            if color.startswith('#'):
                                style_parts.append(f'color: {color};')

                    if cell_val.fill and isinstance(cell_val.fill, PatternFill) and cell_val.fill.start_color:
                        if hasattr(cell_val.fill.start_color, 'rgb') and cell_val.fill.start_color.rgb:
                            bg_color = str(cell_val.fill.start_color.rgb)
                            if len(bg_color) == 8:
                                bg_color = '#' + bg_color[2:]
                            elif len(bg_color) == 6:
                                bg_color = '#' + bg_color
                            if bg_color.startswith('#') and bg_color != '#000000':
                                style_parts.append(f'background-color: {bg_color};')
                    
                    style_attr = f'style="{ " ".join(style_parts) }"'
                    escaped_val = str(value).replace('<', '&lt;').replace('>', '&gt;')
                    escaped_formula = str(formula).replace('"', '&quot;')
                    
                    html_parts.append(f'<td {style_attr} data-formula="{escaped_formula}" onclick="showFormula(this)" style="cursor:pointer;">{escaped_val}</td>')
                
                html_parts.append('</tr>')
            
            html_parts.append('</tbody>')
            html_parts.append('</table>')
            html_parts.append('</div>')
            
            if (sheet_values.max_row or 0) > max_rows:
                html_parts.append(f'<p class="text-muted small mt-2">Показано перші {max_rows} рядків з {sheet_values.max_row}</p>')
                
            html_parts.append('</div>')
        
        html_content = '\n'.join(html_parts)
        return html_content, None

    except Exception as e:
        error_msg = f"Помилка при читанні таблиці: {str(e)}"
        return "", error_msg


def convert_pptx_to_html(file_path: str) -> Tuple[str, Optional[str]]:
    """
    Конвертує .pptx файл у стильні HTML слайди з текстом, зображеннями та таблицями.
    """
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        prs = Presentation(file_path)
        total_slides = len(prs.slides)
        if total_slides == 0:
            return '<p class="text-muted" style="text-align:center;padding:20px;">Презентація не містить слайдів.</p>', None

        html_parts = ['<div class="pptx-viewer">']

        for slide_idx, slide in enumerate(prs.slides, 1):
            html_parts.append('<div class="pptx-slide-card">')
            html_parts.append(
                f'<div class="pptx-slide-header">'
                f'<span class="pptx-slide-num">Слайд {slide_idx} з {total_slides}</span>'
                f'<span class="pptx-slide-badge">📽️ Презентація</span>'
                f'</div>'
            )
            html_parts.append('<div class="pptx-slide-body">')

            # Рекурсивне видобування фігур (включаючи згруповані)
            def extract_shapes(shapes_iterable):
                flat = []
                for shp in shapes_iterable:
                    try:
                        if getattr(shp, 'shape_type', None) == MSO_SHAPE_TYPE.GROUP:
                            flat.extend(extract_shapes(shp.shapes))
                        else:
                            flat.append(shp)
                    except Exception:
                        flat.append(shp)
                return flat

            all_shapes = extract_shapes(slide.shapes)

            pictures = []
            tables = []
            texts = []

            for shape in all_shapes:
                # 1. Зображення
                if getattr(shape, 'shape_type', None) == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, 'image'):
                    try:
                        img_bytes = shape.image.blob
                        img_b64 = base64.b64encode(img_bytes).decode('ascii')
                        mime = getattr(shape.image, 'content_type', 'image/png') or 'image/png'

                        # Обчислюємо реальні розміри фігури на слайді в пікселях (1 px = 9525 EMU)
                        w_px = round(shape.width / 9525) if getattr(shape, 'width', None) else None
                        h_px = round(shape.height / 9525) if getattr(shape, 'height', None) else None

                        # Визначаємо чи це невелика іконка (за розміром фігури на слайді або симетричним квадратним розміром)
                        is_icon = False
                        if w_px and h_px and w_px <= 240 and h_px <= 240:
                            is_icon = True
                        elif hasattr(shape, 'image') and shape.image.size:
                            raw_w, raw_h = shape.image.size
                            if raw_w <= 128 and raw_h <= 128:
                                is_icon = True

                        if is_icon:
                            box_class = "pptx-img-box is-icon"
                            img_class = "pptx-slide-img is-icon"
                            max_w = min(w_px or 120, 140)
                            max_h = min(h_px or 120, 140)
                            style_attr = f'style="max-width:{max_w}px;max-height:{max_h}px;"'
                        elif w_px and h_px:
                            box_class = "pptx-img-box"
                            img_class = "pptx-slide-img"
                            max_w = min(w_px, 860)
                            max_h = min(h_px, 460)
                            style_attr = f'style="max-width:{max_w}px;max-height:{max_h}px;"'
                        else:
                            box_class = "pptx-img-box"
                            img_class = "pptx-slide-img"
                            style_attr = 'style="max-width:100%;max-height:360px;"'

                        pictures.append(
                            f'<div class="{box_class}">'
                            f'<img src="data:{mime};base64,{img_b64}" class="{img_class}" {style_attr} alt="Зображення до слайду {slide_idx}">'
                            f'</div>'
                        )
                    except Exception:
                        pass

                # 2. Таблиці
                elif getattr(shape, 'has_table', False):
                    try:
                        tbl = shape.table
                        t_html = ['<table class="table table-bordered table-striped pptx-slide-table">']
                        for r_idx, row in enumerate(tbl.rows):
                            t_html.append('<tr>')
                            for cell in row.cells:
                                tag = 'th' if r_idx == 0 else 'td'
                                c_text = html.escape(cell.text.strip())
                                t_html.append(f'<{tag}>{c_text}</{tag}>')
                            t_html.append('</tr>')
                        t_html.append('</table>')
                        tables.append(''.join(t_html))
                    except Exception:
                        pass

                # 3. Текстові блоки
                elif hasattr(shape, 'text') and shape.text.strip():
                    texts.append(shape.text.strip())

            content_rendered = False

            # Галерея зображень
            if pictures:
                html_parts.append('<div class="pptx-slide-gallery">')
                for pic in pictures:
                    html_parts.append(pic)
                html_parts.append('</div>')
                content_rendered = True

            # Текст слайду
            if texts:
                html_parts.append('<div class="pptx-slide-text">')
                for t_idx, t in enumerate(texts):
                    # Якщо це схоже на заголовок слайду
                    if t_idx == 0 and len(t) < 120 and '\n' not in t:
                        html_parts.append(f'<h3 class="pptx-slide-title">{html.escape(t)}</h3>')
                    else:
                        for line in t.split('\n'):
                            line_clean = line.strip()
                            if line_clean:
                                if line_clean.startswith(('•', '-', '*', '—')):
                                    html_parts.append(f'<div class="pptx-bullet">• {html.escape(line_clean.lstrip("•-*— "))}</div>')
                                else:
                                    html_parts.append(f'<p>{html.escape(line_clean)}</p>')
                html_parts.append('</div>')
                content_rendered = True

            # Таблиці слайду
            if tables:
                for tbl_html in tables:
                    html_parts.append(f'<div class="pptx-slide-table-wrapper">{tbl_html}</div>')
                content_rendered = True

            if not content_rendered:
                html_parts.append('<p class="text-muted" style="text-align:center;padding:24px 0;">Слайд без текстового чи візуального контенту</p>')

            html_parts.append('</div>')  # .pptx-slide-body
            html_parts.append('</div>')  # .pptx-slide-card

        html_parts.append('</div>')  # .pptx-viewer
        return '\n'.join(html_parts), None

    except Exception as e:
        error_msg = f"Помилка при читанні презентації: {str(e)}"
        return "", error_msg


def convert_odt_to_html(file_path: str) -> Tuple[str, Optional[str]]:
    """Конвертує .odt файл у HTML."""
    try:
        from odf.opendocument import load
        from odf.text import P, H
        from odf import teletype

        doc = load(file_path)
        html_parts = ['<div class="document-page">']

        for h in doc.getElementsByType(H):
            text = teletype.extractText(h)
            level = h.getAttribute('outlinelevel') or 3
            html_parts.append(f'<h{level}>{text}</h{level}>')

        for p in doc.getElementsByType(P):
            text = teletype.extractText(p)
            if text.strip():
                html_parts.append(f'<p>{text}</p>')

        html_parts.append('</div>')
        html_content = '\n'.join(html_parts)
        
        if not html_content.strip() or html_content == '<div class="document-page"></div>':
            html_content = '<p class="text-muted">Документ порожній або не містить текстового контенту.</p>'
            
        return html_content, None

    except Exception as e:
        error_msg = f"Помилка при читанні ODT документа: {str(e)}"
        return "", error_msg


def convert_ods_to_html(file_path: str, max_rows: int = 100) -> Tuple[str, Optional[str]]:
    """Конвертує .ods файл у HTML таблиці."""
    try:
        from odf.opendocument import load
        from odf.table import Table, TableRow, TableCell
        from odf import teletype

        doc = load(file_path)
        html_parts = ['<div class="document-page">']

        for sheet in doc.getElementsByType(Table):
            sheet_name = sheet.getAttribute('name')
            html_parts.append(f'<h4 class="mb-3" style="color:var(--color-primary,#6366f1);font-weight:700;">Аркуш: {sheet_name}</h4>')
            html_parts.append('<div class="excel-sheet mb-4 table-responsive">')
            html_parts.append('<table class="excel-table table table-bordered" style="font-size:13px;">')
            
            rows = sheet.getElementsByType(TableRow)
            for i, row in enumerate(rows):
                if i >= max_rows:
                    break
                    
                html_parts.append('<tr>')
                html_parts.append(f'<td class="row-header" style="background:var(--color-bg-secondary,#f8fafc);text-align:center;font-weight:600;width:40px;">{i+1}</td>')
                
                cells = row.getElementsByType(TableCell)
                for cell in cells:
                    repeat = int(cell.getAttribute('numbercolumnsrepeated') or 1)
                    text = teletype.extractText(cell)
                    formula = cell.getAttribute('formula') or ""
                    
                    if formula.startswith('of:='):
                        formula = '=' + formula[4:]
                    elif not formula and text:
                        formula = text
                        
                    for _ in range(repeat):
                        escaped_val = str(text).replace('<', '&lt;').replace('>', '&gt;')
                        escaped_form = str(formula).replace('"', '&quot;')
                        html_parts.append(f'<td data-formula="{escaped_form}" onclick="showFormula(this)" style="cursor:pointer;padding:6px 10px;">{escaped_val}</td>')
                        
                html_parts.append('</tr>')
            
            html_parts.append('</table>')
            html_parts.append('</div>')
            
            if len(rows) > max_rows:
                html_parts.append(f'<p class="text-muted small">Показано перші {max_rows} рядків з {len(rows)}</p>')

        html_parts.append('</div>')
        html_content = '\n'.join(html_parts)
        return html_content, None

    except Exception as e:
        error_msg = f"Помилка при читанні ODS таблиці: {str(e)}"
        return "", error_msg


def convert_odp_to_html(file_path: str) -> Tuple[str, Optional[str]]:
    """Конвертує .odp файл у HTML."""
    try:
        from odf.opendocument import load
        from odf.draw import Page, Frame, TextBox
        from odf import teletype

        doc = load(file_path)
        html_parts = ['<div class="pptx-viewer" style="display:flex;flex-direction:column;gap:20px;">']

        slides = doc.getElementsByType(Page)
        for i, slide in enumerate(slides, 1):
            html_parts.append(f'<div class="slide-container card p-4 shadow-sm" style="border:1px solid var(--color-border,#e2e8f0);border-radius:12px;background:var(--color-surface,#fff);">')
            html_parts.append(f'<div class="d-flex align-items-center justify-content-between mb-3 pb-2 border-bottom">')
            html_parts.append(f'<h5 style="margin:0;color:var(--color-primary,#6366f1);font-weight:700;">Слайд {i}</h5>')
            html_parts.append('</div>')
            html_parts.append('<div class="slide-body" style="font-size:15px;line-height:1.6;">')
            
            slide_text_found = False
            for frame in slide.getElementsByType(Frame):
                for textbox in frame.getElementsByType(TextBox):
                    text = teletype.extractText(textbox)
                    if text.strip():
                        html_parts.append(f'<p style="margin-bottom:8px;">{text}</p>')
                        slide_text_found = True
            
            if not slide_text_found:
                html_parts.append('<p class="text-muted">Слайд не містить текстового контенту</p>')
                
            html_parts.append('</div>')
            html_parts.append('</div>')

        html_parts.append('</div>')
        html_content = '\n'.join(html_parts)
        
        if not slides:
            html_content = '<p class="text-muted">Презентація не містить слайдів.</p>'

        return html_content, None

    except Exception as e:
        error_msg = f"Помилка при читанні ODP презентації: {str(e)}"
        return "", error_msg


def get_archive_content(file_path: str, file_ext: str) -> Tuple[list, Optional[str]]:
    """Повертає список файлів в архіві (.zip, .rar, .7z, .tar, .gz)."""
    files_list = []
    try:
        if file_ext == '.zip':
            import zipfile
            with zipfile.ZipFile(file_path, 'r') as zf:
                for info in zf.infolist():
                    files_list.append({
                        'name': info.filename,
                        'size': info.file_size,
                        'is_dir': info.is_dir()
                    })
                    
        elif file_ext == '.rar':
            try:
                import rarfile
                with rarfile.RarFile(file_path) as rf:
                    for info in rf.infolist():
                        files_list.append({
                            'name': info.filename,
                            'size': info.file_size,
                            'is_dir': info.isdir()
                        })
            except Exception:
                pass
                
        elif file_ext == '.7z':
            try:
                import py7zr
                with py7zr.SevenZipFile(file_path, 'r') as zf:
                    for info in zf.list():
                        files_list.append({
                            'name': info.filename,
                            'size': info.uncompressed,
                            'is_dir': info.is_directory
                        })
            except Exception:
                pass

        elif file_ext in ('.tar', '.gz', '.tgz'):
            import tarfile
            with tarfile.open(file_path, 'r:*') as tf:
                for member in tf.getmembers():
                    files_list.append({
                        'name': member.name,
                        'size': member.size,
                        'is_dir': member.isdir()
                    })
                    
        return files_list, None
        
    except Exception as e:
        error_msg = f"Помилка при читанні архіву: {str(e)}"
        return [], error_msg


def get_file_type_info(file_ext: str) -> dict:
    """Повертає інформацію про тип файлу (назва, іконка, прев'ю)."""
    file_types = {
        '.doc': {'name': 'Word документ', 'icon': '📝', 'color': 'primary', 'preview': False},
        '.docx': {'name': 'Word документ', 'icon': '📝', 'color': 'primary', 'preview': True},
        '.odt': {'name': 'OpenDocument Текст', 'icon': '📝', 'color': 'primary', 'preview': True},
        
        '.xls': {'name': 'Excel таблиця', 'icon': '📊', 'color': 'success', 'preview': False},
        '.xlsx': {'name': 'Excel таблиця', 'icon': '📊', 'color': 'success', 'preview': True},
        '.ods': {'name': 'OpenDocument Таблиця', 'icon': '📊', 'color': 'success', 'preview': True},
        
        '.ppt': {'name': 'PowerPoint', 'icon': '📽️', 'color': 'warning', 'preview': False},
        '.pptx': {'name': 'PowerPoint', 'icon': '📽️', 'color': 'warning', 'preview': True},
        '.odp': {'name': 'OpenDocument Презентація', 'icon': '📽️', 'color': 'warning', 'preview': True},
        
        '.pdf': {'name': 'PDF документ', 'icon': '📄', 'color': 'danger', 'preview': True},
        
        '.jpg': {'name': 'Зображення', 'icon': '🖼️', 'color': 'info', 'preview': True},
        '.jpeg': {'name': 'Зображення', 'icon': '🖼️', 'color': 'info', 'preview': True},
        '.png': {'name': 'Зображення', 'icon': '🖼️', 'color': 'info', 'preview': True},
        '.gif': {'name': 'Зображення', 'icon': '🖼️', 'color': 'info', 'preview': True},
        '.bmp': {'name': 'Зображення', 'icon': '🖼️', 'color': 'info', 'preview': True},
        '.webp': {'name': 'Зображення', 'icon': '🖼️', 'color': 'info', 'preview': True},
        
        '.txt': {'name': 'Текстовий файл', 'icon': '📃', 'color': 'secondary', 'preview': True},
        '.md': {'name': 'Markdown документ', 'icon': '📃', 'color': 'secondary', 'preview': True},
        
        '.py': {'name': 'Python', 'icon': '🐍', 'color': 'teal', 'preview': True},
        '.html': {'name': 'HTML', 'icon': '💻', 'color': 'teal', 'preview': True},
        '.css': {'name': 'CSS', 'icon': '💻', 'color': 'teal', 'preview': True},
        '.js': {'name': 'JavaScript', 'icon': '💻', 'color': 'teal', 'preview': True},
        '.json': {'name': 'JSON', 'icon': '💻', 'color': 'teal', 'preview': True},
        '.xml': {'name': 'XML', 'icon': '💻', 'color': 'teal', 'preview': True},
        
        '.zip': {'name': 'ZIP архів', 'icon': '📦', 'color': 'secondary', 'preview': True},
        '.rar': {'name': 'RAR архів', 'icon': '📦', 'color': 'secondary', 'preview': True},
        '.7z': {'name': '7Z архів', 'icon': '📦', 'color': 'secondary', 'preview': True},
        '.tar': {'name': 'TAR архів', 'icon': '📦', 'color': 'secondary', 'preview': True},
        '.gz': {'name': 'GZ архів', 'icon': '📦', 'color': 'secondary', 'preview': True},

        # Microsoft Access Databases
        '.mdb': {'name': 'Microsoft Access 97-2003', 'icon': '🗄️', 'color': 'success', 'preview': True},
        '.accdb': {'name': 'Microsoft Access 2007+', 'icon': '🗄️', 'color': 'success', 'preview': True},
    }
    
    return file_types.get(file_ext.lower() if file_ext else '', {
        'name': f'{(file_ext or "").upper().replace(".", "")} файл',
        'icon': '📎',
        'color': 'secondary',
        'preview': False
    })


def parse_embed_url(url: str) -> dict:
    """
    Аналізує посилання та перетворює його на валідну embed-URL для <iframe>.
    """
    if not url:
        return {
            'is_embeddable': False,
            'embed_type': None,
            'embed_url': None,
            'original_url': '',
            'platform': None,
            'icon': '🔗',
            'video_id': None
        }

    clean_url = url.strip()

    # 1. YouTube
    yt_regex = (
        r'(?:https?:\/\/)?(?:www\.|m\.)?'
        r'(?:youtube\.com\/(?:watch\?(?:.*&)?v=|embed\/|v\/|shorts\/|live\/)|youtu\.be\/)'
        r'([a-zA-Z0-9_-]{11})'
    )
    yt_match = re.search(yt_regex, clean_url)
    if yt_match:
        video_id = yt_match.group(1)
        parsed = urlparse(clean_url)
        params = parse_qs(parsed.query)
        start_param = ''
        if 't' in params:
            t_str = params['t'][0].replace('s', '')
            if t_str.isdigit():
                start_param = f"?start={t_str}"
        elif 'start' in params:
            if params['start'][0].isdigit():
                start_param = f"?start={params['start'][0]}"

        embed_url = f"https://www.youtube.com/embed/{video_id}{start_param}"
        return {
            'is_embeddable': True,
            'embed_type': 'youtube',
            'embed_url': embed_url,
            'original_url': clean_url,
            'platform': 'YouTube',
            'icon': '▶️',
            'video_id': video_id
        }

    # 2. Vimeo
    vimeo_regex = r'(?:https?:\/\/)?(?:www\.)?vimeo\.com\/(?:channels\/(?:\w+\/)?|groups\/([^\/]*)\/videos\/|album\/(\d+)\/video\/|)(\d+)'
    vimeo_match = re.search(vimeo_regex, clean_url)
    if vimeo_match:
        video_id = vimeo_match.group(3)
        return {
            'is_embeddable': True,
            'embed_type': 'vimeo',
            'embed_url': f"https://player.vimeo.com/video/{video_id}",
            'original_url': clean_url,
            'platform': 'Vimeo',
            'icon': '▶️',
            'video_id': video_id
        }

    # 3. Google Drive
    if 'drive.google.com' in clean_url and '/file/d/' in clean_url:
        gdrive_match = re.search(r'/file/d/([a-zA-Z0-9_-]+)', clean_url)
        if gdrive_match:
            gdrive_id = gdrive_match.group(1)
            return {
                'is_embeddable': True,
                'embed_type': 'google_drive',
                'embed_url': f"https://drive.google.com/file/d/{gdrive_id}/preview",
                'original_url': clean_url,
                'platform': 'Google Диск',
                'icon': '📁',
                'video_id': gdrive_id
            }

    # 4. Google Docs / Sheets / Slides
    if 'docs.google.com' in clean_url:
        doc_preview = re.sub(r'/(edit|view|copy|export).*$', '/preview', clean_url)
        if not doc_preview.endswith('/preview'):
            doc_preview = doc_preview.rstrip('/') + '/preview'
        return {
            'is_embeddable': True,
            'embed_type': 'google_docs',
            'embed_url': doc_preview,
            'original_url': clean_url,
            'platform': 'Google Документи',
            'icon': '📑',
            'video_id': None
        }

    # Generic Website / Link
    return {
        'is_embeddable': True,
        'embed_type': 'generic',
        'embed_url': clean_url,
        'original_url': clean_url,
        'platform': 'Веб-посилання',
        'icon': '🔗',
        'video_id': None
    }


def optimize_uploaded_file(uploaded_file):
    """
    Оптимізує завантажений файл: якщо це зображення (JPG, PNG тощо),
    виправляє EXIF-орієнтацію, зменшує роздільну здатність до макс. 1920px
    та зберігає у легкому сучасному форматі WebP для економії пам'яті сервера.
    """
    if not uploaded_file:
        return uploaded_file

    ext = os.path.splitext(uploaded_file.name)[1].lower()
    if ext not in {'.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.webp'}:
        return uploaded_file

    try:
        from io import BytesIO
        from PIL import Image, ImageOps
        from django.core.files.uploadedfile import InMemoryUploadedFile

        img = Image.open(uploaded_file)
        # Виправлення орієнтації за EXIF
        img = ImageOps.exif_transpose(img)

        # Конвертація кольорової моделі
        if img.mode in ('RGBA', 'LA', 'P'):
            img = img.convert('RGBA')
        else:
            img = img.convert('RGB')

        # Зменшення якщо роздільна здатність перевищує 1920px
        max_dim = 1920
        if max(img.size) > max_dim:
            img.thumbnail((max_dim, max_dim), Image.Resampling.LANCZOS)

        output = BytesIO()
        img.save(output, format='WEBP', quality=85, optimize=True)
        output.seek(0)

        base_name = os.path.splitext(uploaded_file.name)[0]
        new_filename = f"{base_name}.webp"

        optimized_file = InMemoryUploadedFile(
            file=output,
            field_name=getattr(uploaded_file, 'field_name', 'file'),
            name=new_filename,
            content_type='image/webp',
            size=output.getbuffer().nbytes,
            charset=None
        )
        return optimized_file
    except Exception:
        return uploaded_file


# ═══════════════════════════════════════════════════════════════════════════════
# РОЗКЛАД УРОКІВ ТА ЦЕНТР СПОВІЩЕНЬ
# ═══════════════════════════════════════════════════════════════════════════════

def get_teacher_live_lesson_status(teacher, now_dt=None):
    """
    Обчислює стан поточного уроку / перерви для вчителя в реальному часі.
    Повертає словник зі статусом, назвою уроку, залишковим часом тощо.
    """
    from django.utils import timezone
    from .models import TeacherLessonSchedule

    if not teacher or not teacher.lesson_schedules.exists():
        return {
            'has_schedule': False,
            'status_type': 'no_schedule',
            'title': 'Розклад не налаштовано',
            'time_info': 'Ви можете додати розклад уроків у блоці «Учні та класи»',
            'badge_text': 'Розклад не налаштовано',
            'badge_class': 'badge-secondary',
            'current_lesson': None,
            'next_lesson': None,
            'remaining_minutes': 0,
        }

    now = now_dt or timezone.localtime(timezone.now())
    today_date = now.date()
    today_weekday = today_date.weekday() + 1  # 1=Понеділок ... 7=Неділя
    current_time = now.time()
    now_minutes = current_time.hour * 60 + current_time.minute

    # Уроки вчителя на сьогодні
    today_lessons = list(
        TeacherLessonSchedule.objects.filter(teacher=teacher, day_of_week=today_weekday)
        .select_related('bell_slot', 'class_group', 'subject')
        .order_by('bell_slot__lesson_number')
    )

    uk_weekdays = {
        1: 'Понеділок',
        2: 'Вівторок',
        3: 'Середа',
        4: 'Четвер',
        5: "П'ятниця",
        6: 'Субота',
        7: 'Неділя',
    }

    if not today_lessons:
        future_lessons = TeacherLessonSchedule.objects.filter(
            teacher=teacher,
            day_of_week__gt=today_weekday
        ).select_related('bell_slot', 'class_group', 'subject').order_by('day_of_week', 'bell_slot__lesson_number')
        next_les = future_lessons.first()
        if not next_les:
            next_les = TeacherLessonSchedule.objects.filter(
                teacher=teacher
            ).select_related('bell_slot', 'class_group', 'subject').order_by('day_of_week', 'bell_slot__lesson_number').first()

        next_info = ""
        if next_les:
            day_name = uk_weekdays.get(next_les.day_of_week, '')
            start_str = next_les.bell_slot.start_time.strftime('%H:%M') if next_les.bell_slot else ''
            next_info = f"Наступний урок: {day_name}, {next_les.bell_slot.lesson_number}-й ({next_les.class_group.name}) о {start_str}"

        return {
            'has_schedule': True,
            'status_type': 'no_lessons_today',
            'title': 'Сьогодні уроків немає',
            'time_info': next_info or 'Відпочивайте!',
            'badge_text': '☕ Без уроків сьогодні',
            'badge_class': 'badge-secondary',
            'current_lesson': None,
            'next_lesson': next_les,
            'remaining_minutes': 0,
        }

    first_lesson = today_lessons[0]
    last_lesson = today_lessons[-1]

    first_start_min = first_lesson.bell_slot.start_time.hour * 60 + first_lesson.bell_slot.start_time.minute

    # 1. До початку першого уроку
    if now_minutes < first_start_min:
        diff = first_start_min - now_minutes
        start_str = first_lesson.bell_slot.start_time.strftime('%H:%M')
        subj_str = f" ({first_lesson.subject.name})" if first_lesson.subject else ""
        return {
            'has_schedule': True,
            'status_type': 'before_school',
            'title': f"1-й за розкладом: {first_lesson.bell_slot.lesson_number}-й урок • {first_lesson.class_group.name}{subj_str}",
            'time_info': f"до початку уроку залишилось {diff} хв (початок о {start_str})",
            'badge_text': '🌅 Скоро початок',
            'badge_class': 'badge-info',
            'current_lesson': None,
            'next_lesson': first_lesson,
            'remaining_minutes': diff,
        }

    # 2. Перевіряємо кожен урок та перерви між ними
    for i, lesson in enumerate(today_lessons):
        slot = lesson.bell_slot
        s_min = slot.start_time.hour * 60 + slot.start_time.minute
        e_min = slot.end_time.hour * 60 + slot.end_time.minute
        subj_str = f" ({lesson.subject.name})" if lesson.subject else ""

        # Урок триває зараз
        if s_min <= now_minutes < e_min:
            rem = e_min - now_minutes
            next_l = today_lessons[i + 1] if i + 1 < len(today_lessons) else None
            end_str = slot.end_time.strftime('%H:%M')
            return {
                'has_schedule': True,
                'status_type': 'in_lesson',
                'title': f"Зараз {slot.lesson_number}-й урок: {lesson.class_group.name}{subj_str}",
                'time_info': f"до кінця уроку залишилось {rem} хв (до {end_str})",
                'badge_text': '🟢 Зараз триває урок',
                'badge_class': 'badge-success',
                'current_lesson': lesson,
                'next_lesson': next_l,
                'remaining_minutes': rem,
            }

        # Перерва між уроками
        if i + 1 < len(today_lessons):
            next_l = today_lessons[i + 1]
            next_s_min = next_l.bell_slot.start_time.hour * 60 + next_l.bell_slot.start_time.minute
            if e_min <= now_minutes < next_s_min:
                rem = next_s_min - now_minutes
                total_break = next_s_min - e_min
                next_start_str = next_l.bell_slot.start_time.strftime('%H:%M')
                next_subj = f" ({next_l.subject.name})" if next_l.subject else ""
                return {
                    'has_schedule': True,
                    'status_type': 'in_break',
                    'title': f"Зараз перерва ({total_break} хв)",
                    'time_info': f"до початку {next_l.bell_slot.lesson_number}-го уроку ({next_l.class_group.name}{next_subj}) залишилось {rem} хв (о {next_start_str})",
                    'badge_text': '☕ Перерва',
                    'badge_class': 'badge-warning',
                    'current_lesson': None,
                    'next_lesson': next_l,
                    'remaining_minutes': rem,
                }

    # 3. Після останнього уроку сьогодні
    return {
        'has_schedule': True,
        'status_type': 'after_school',
        'title': 'Уроки на сьогодні завершено',
        'time_info': f"Останній урок закінчився о {last_lesson.bell_slot.end_time.strftime('%H:%M')}",
        'badge_text': '🏁 Заняття закінчено',
        'badge_class': 'badge-secondary',
        'current_lesson': None,
        'next_lesson': None,
        'remaining_minutes': 0,
    }


def get_teacher_upcoming_notifications(teacher, now_dt=None):
    """
    Формує список активних сповіщень для вчителя:
    1. Урок починається скоро або триває прямо зараз, але завдання для цього класу не опубліковано.
    2. Нові неперевірені здані роботи від учнів.
    """
    from django.utils import timezone
    from django.urls import reverse
    from .models import TeacherLessonSchedule, Assignment, Submission

    notifications = []
    if not teacher:
        return notifications

    now = now_dt or timezone.localtime(timezone.now())
    today = now.date()
    today_weekday = today.weekday() + 1
    now_min = now.time().hour * 60 + now.time().minute

    # 1. Перевіряємо уроки на сьогодні
    today_lessons = TeacherLessonSchedule.objects.filter(
        teacher=teacher,
        day_of_week=today_weekday
    ).select_related('bell_slot', 'class_group', 'subject').order_by('bell_slot__lesson_number')

    for lesson in today_lessons:
        slot = lesson.bell_slot
        s_min = slot.start_time.hour * 60 + slot.start_time.minute
        e_min = slot.end_time.hour * 60 + slot.end_time.minute

        # Якщо урок розпочнеться в найближчі 45 хвилин АБО триває прямо зараз
        is_upcoming_or_now = (s_min - 45 <= now_min < e_min)
        if is_upcoming_or_now:
            # Шукаємо, чи є опубліковане завдання для цього класу на сьогодні/для цього уроку
            has_task = Assignment.objects.filter(
                teacher=teacher,
                classes=lesson.class_group,
                status=Assignment.STATUS_PUBLISHED
            ).filter(
                published_at__date=today
            ).exists()

            if not has_task:
                has_targeted_task = Assignment.objects.filter(
                    teacher=teacher,
                    classes=lesson.class_group,
                    status=Assignment.STATUS_PUBLISHED,
                    schedule_targets__class_group=lesson.class_group,
                    schedule_targets__target_date=today
                ).exists()
                if has_targeted_task:
                    has_task = True

            if not has_task:
                status_text = "зараз триває" if (s_min <= now_min < e_min) else f"почнеться о {slot.start_time.strftime('%H:%M')}"
                create_url = f"{reverse('assignment_create')}?class={lesson.class_group_id}&subject={lesson.subject_id or ''}"
                notifications.append({
                    'id': f"missing_task_{lesson.id}_{today}",
                    'type': 'missing_task',
                    'icon': '⚠️',
                    'title': f"Немає завдання для {lesson.class_group.name}",
                    'message': f"{slot.lesson_number}-й урок {status_text}, але для {lesson.class_group.name} ще не опубліковано завдання!",
                    'action_url': create_url,
                    'action_label': f"➕ Створити завдання для {lesson.class_group.name}",
                    'badge': 'Терміново',
                    'badge_class': 'badge-danger',
                    'time_str': slot.start_time.strftime('%H:%M'),
                })

    # 2. Неперевірені здані роботи
    ungraded_count = Submission.objects.filter(
        assignment__teacher=teacher,
        grade__isnull=True
    ).count() + Submission.objects.filter(
        assignment__teacher=teacher,
        grade=''
    ).count()

    if ungraded_count > 0:
        notifications.append({
            'id': 'pending_submissions',
            'type': 'pending_submissions',
            'icon': '📥',
            'title': 'Здані роботи учнів',
            'message': f"Очікують вашої перевірки та оцінювання: {ungraded_count} робіт.",
            'action_url': reverse('all_submissions_dashboard'),
            'action_label': 'Перейти до робіт ↗',
            'badge': str(ungraded_count),
            'badge_class': 'badge-warning',
            'time_str': 'Сьогодні',
        })

    return notifications


