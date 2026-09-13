"""
Скрипт повної ініціалізації тестового середовища SchoolNet+ для користувача.
Створює:
1. Адміністратора з логіном та паролем.
2. Профіль вчителя з усіма класами та предметами.
3. Розклад дзвінків (BellSchedule) на 7 уроків.
4. Розклад уроків вчителя (TeacherLessonSchedule) на пн-пт.
5. Різні типи завдань:
   - Завдання для кількох класів з різними датами та часом (Вимога 5)
   - Завдання з форматованим текстом (Вимога 8)
   - Завдання з прикріпленим DOCX-файлом (Вимога 6)
   - Чернетка (Вимоги 2 та 4)
   - Здачі робіт учнями для перевірки центру сповіщень та масового оцінювання (Вимоги 1 та 7)
"""

import os
import sys
import io
import datetime
import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'schoolnet.settings')
django.setup()

from django.contrib.auth.models import User
from django.utils import timezone
from django.core.files.uploadedfile import SimpleUploadedFile
import docx

from feed.models import (
    Teacher, ClassGroup, Subject, BellSchedule, TeacherLessonSchedule,
    Assignment, AssignmentFile, AssignmentLink, AssignmentScheduleTarget,
    Submission, SubmissionComment, AICriteriaPreset
)

print("🌟 [1/6] Створення/оновлення облікового запису адміністратора...")
ADMIN_USERNAME = "admin"
ADMIN_PASSWORD = "SchoolNet2026!"
ADMIN_FULL_NAME = "Олександр Васильович Коваленко"
ADMIN_EMAIL = "admin@schoolnet.ua"

admin_user, created = User.objects.get_or_create(username=ADMIN_USERNAME)
admin_user.set_password(ADMIN_PASSWORD)
admin_user.is_staff = True
admin_user.is_superuser = True
admin_user.email = ADMIN_EMAIL
admin_user.first_name = "Олександр"
admin_user.last_name = "Коваленко"
admin_user.save()

# Також оновить пароль для yrsadmin, якщо користувач звик до нього
yrs_user = User.objects.filter(username='yrsadmin').first()
if yrs_user:
    yrs_user.set_password(ADMIN_PASSWORD)
    yrs_user.is_staff = True
    yrs_user.is_superuser = True
    yrs_user.save()

teacher, _ = Teacher.objects.get_or_create(
    user=admin_user,
    defaults={'full_name': ADMIN_FULL_NAME}
)
teacher.full_name = ADMIN_FULL_NAME
teacher.avatar_color = "#4f46e5"
teacher.save()

# Ініціалізація НУШ критеріїв
try:
    AICriteriaPreset.ensure_default_presets()
except Exception:
    pass

print(f"  ✅ Адміністратор готовий: логін '{ADMIN_USERNAME}', пароль '{ADMIN_PASSWORD}'")

print("🏫 [2/6] Налаштування класів та предметів...")
classes_data = [
    ('5А', 5, 'А'), ('6А', 6, 'А'), ('7А', 7, 'А'),
    ('8А', 8, 'А'), ('9А', 9, 'А'), ('9Б', 9, 'Б'),
    ('10А', 10, 'А'), ('11А', 11, 'А')
]
class_objs = {}
for name, grade, letter in classes_data:
    c, _ = ClassGroup.objects.get_or_create(name=name, defaults={'grade': grade, 'letter': letter})
    class_objs[name] = c
    teacher.classes.add(c)

subj_data = [
    ('Інформатика', '💻', '#3b82f6'),
    ('Математика', '📐', '#10b981'),
    ('Фізика', '⚡', '#8b5cf6'),
    ('Українська мова', '📖', '#f59e0b'),
    ('Англійська мова', '🌍', '#ec4899')
]
subj_objs = {}
for name, icon, col in subj_data:
    s, _ = Subject.objects.get_or_create(name=name, defaults={'icon': icon, 'color': col})
    subj_objs[name] = s
    teacher.subjects.add(s)

print("⏰ [3/6] Створення розкладу дзвінків...")
bells_data = [
    (1, datetime.time(8, 30), datetime.time(9, 15)),
    (2, datetime.time(9, 25), datetime.time(10, 10)),
    (3, datetime.time(10, 25), datetime.time(11, 10)),
    (4, datetime.time(11, 25), datetime.time(12, 10)),
    (5, datetime.time(12, 25), datetime.time(13, 10)),
    (6, datetime.time(13, 20), datetime.time(14, 5)),
    (7, datetime.time(14, 15), datetime.time(15, 0)),
]
bell_objs = {}
for num, st, et in bells_data:
    b, _ = BellSchedule.objects.get_or_create(lesson_number=num, defaults={'start_time': st, 'end_time': et, 'order': num})
    bell_objs[num] = b

print("📅 [4/6] Створення розкладу уроків вчителя на тиждень...")
TeacherLessonSchedule.objects.filter(teacher=teacher).delete()

schedule_rows = [
    # Понеділок (1)
    (1, 2, '7А', 'Інформатика'),
    (1, 3, '8А', 'Інформатика'),
    (1, 4, '9А', 'Інформатика'),
    # Вівторок (2)
    (2, 1, '10А', 'Інформатика'),
    (2, 2, '9Б', 'Інформатика'),
    (2, 3, '7А', 'Математика'),
    # Середа (3)
    (3, 2, '8А', 'Математика'),
    (3, 3, '9А', 'Математика'),
    (3, 4, '11А', 'Інформатика'),
    # Четвер (4)
    (4, 1, '7А', 'Інформатика'),
    (4, 2, '8А', 'Інформатика'),
    (4, 3, '9Б', 'Інформатика'),
    # П'ятниця (5)
    (5, 2, '10А', 'Інформатика'),
    (5, 3, '11А', 'Інформатика'),
]

for dow, slot_num, cls_name, sub_name in schedule_rows:
    TeacherLessonSchedule.objects.create(
        teacher=teacher,
        class_group=class_objs[cls_name],
        subject=subj_objs[sub_name],
        day_of_week=dow,
        bell_slot=bell_objs[slot_num]
    )

print("📝 [5/6] Створення тестових завдань...")
now = timezone.localtime(timezone.now())
today = now.date()
# Визначаємо дати на поточному тижні
monday = today - datetime.timedelta(days=today.weekday())
tuesday = monday + datetime.timedelta(days=1)
wednesday = monday + datetime.timedelta(days=2)
thursday = monday + datetime.timedelta(days=3)
friday = monday + datetime.timedelta(days=4)

# 1. Завдання для декількох класів з РІЗНИМ розкладом та часом
task_multiclass = Assignment.objects.create(
    teacher=teacher,
    subject=subj_objs['Інформатика'],
    title='Практична робота: Основи мови Python та структури даних',
    description="""<h3>Мета роботи</h3>
<p>Ознайомитися з синтаксисом <b>Python</b>, операторами розгалуження <code>if-else</code> та списками.</p>
<h4>Вимоги до виконання:</h4>
<ul>
  <li>Створити скрипт <u>solution.py</u></li>
  <li>Реалізувати введення чисел з консолі</li>
  <li>Вивести відсортований список результатів</li>
</ul>
<p><i>Зверніть увагу: код повинен містити коментарі до ключових функцій.</i></p>""",
    status=Assignment.STATUS_PUBLISHED,
    published_at=now - datetime.timedelta(days=1),
)
task_multiclass.classes.add(class_objs['7А'], class_objs['8А'], class_objs['9А'])

AssignmentScheduleTarget.objects.create(
    assignment=task_multiclass,
    class_group=class_objs['7А'],
    target_day_of_week=1,
    bell_slot=bell_objs[2], # 09:25
    target_date=monday
)
AssignmentScheduleTarget.objects.create(
    assignment=task_multiclass,
    class_group=class_objs['8А'],
    target_day_of_week=1,
    bell_slot=bell_objs[3], # 10:25
    target_date=monday
)
AssignmentScheduleTarget.objects.create(
    assignment=task_multiclass,
    class_group=class_objs['9А'],
    target_day_of_week=1,
    bell_slot=bell_objs[4], # 11:25
    target_date=monday
)

# 2. Завдання з прикріпленим реальним DOCX-документом (для тестування DOCX-переглядача)
doc = docx.Document()
doc.add_heading('Інструкційна картка до лабораторної роботи №3', level=1)
doc.add_paragraph('Тема: Комп\'ютерні мережі та безпека даних у локальній мережі школи.')
p = doc.add_paragraph()
p.add_run('Завдання 1: ').bold = True
p.add_run('Дослідити IP-конфігурацію власного робочого місця за допомогою команди ipconfig / ifconfig.')
table = doc.add_table(rows=3, cols=3)
table.style = 'Table Grid'
table.cell(0, 0).text = 'Параметр'
table.cell(0, 1).text = 'Значення'
table.cell(0, 2).text = 'Призначення'
table.cell(1, 0).text = 'IP-адреса'
table.cell(1, 1).text = '192.168.1.105'
table.cell(1, 2).text = 'Вузол мережі'
table.cell(2, 0).text = 'Маска підмережі'
table.cell(2, 1).text = '255.255.255.0'
table.cell(2, 2).text = 'Клас C'

doc_io = io.BytesIO()
doc.save(doc_io)
doc_io.seek(0)

uploaded_docx = SimpleUploadedFile(
    'Instruktsiyna_kartka_merezhi.docx',
    doc_io.read(),
    content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
)

task_docx = Assignment.objects.create(
    teacher=teacher,
    subject=subj_objs['Інформатика'],
    title='Лабораторна робота №3: Комп\'ютерні мережі (DOCX матеріали)',
    description='<p>Ознайомтеся з вкладеною інструкцією у форматі <b>Word (DOCX)</b> та заповніть звіт.</p>',
    status=Assignment.STATUS_PUBLISHED,
    published_at=now,
)
task_docx.classes.add(class_objs['9Б'])
AssignmentScheduleTarget.objects.create(
    assignment=task_docx,
    class_group=class_objs['9Б'],
    target_day_of_week=2,
    bell_slot=bell_objs[2],
    target_date=tuesday
)

AssignmentFile.objects.create(
    assignment=task_docx,
    file=uploaded_docx,
    original_name='Instruktsiyna_kartka_merezhi.docx'
)

# 3. Чернетка (Draft) — не відображається учням і не рахується в проведених уроках
task_draft = Assignment.objects.create(
    teacher=teacher,
    subject=subj_objs['Математика'],
    title='[Чернетка] Контрольна робота: Квадратні рівняння',
    description='<p>Чернетка варіантів контрольної роботи до узгодження з методистом.</p>',
    status=Assignment.STATUS_DRAFT,
)
task_draft.classes.add(class_objs['8А'])
AssignmentScheduleTarget.objects.create(
    assignment=task_draft,
    class_group=class_objs['8А'],
    target_day_of_week=3,
    bell_slot=bell_objs[2],
    target_date=wednesday
)

# 4. Завдання на наступний понеділок (для тестування відображення в календарі)
next_monday = monday + datetime.timedelta(days=7)
task_calendar = Assignment.objects.create(
    teacher=teacher,
    subject=subj_objs['Інформатика'],
    title='Завдання на наступний тиждень: Алгоритми сортування',
    description='<p>Підготувати конспект про алгоритми сортування бульбашкою та швидкого сортування (Quicksort).</p>',
    status=Assignment.STATUS_PUBLISHED,
    published_at=now,
)
task_calendar.classes.add(class_objs['10А'])
AssignmentScheduleTarget.objects.create(
    assignment=task_calendar,
    class_group=class_objs['10А'],
    target_day_of_week=1,
    bell_slot=bell_objs[1],
    target_date=next_monday
)

print("👥 [6/6] Створення учнівських здач (робіт) для тестування оцінювання...")
# Створюємо кілька зданих робіт
sub1 = Submission.objects.create(
    assignment=task_multiclass,
    first_name='Максим',
    last_name='Бондаренко',
    class_group=class_objs['7А'],
    teacher=teacher,
    comment_student='Виконав усі пункти завдання, програма протестована на Python 3.12.',
    submitted_at=now - datetime.timedelta(hours=2)
)

sub2 = Submission.objects.create(
    assignment=task_multiclass,
    first_name='Софія',
    last_name='Мельник',
    class_group=class_objs['7А'],
    teacher=teacher,
    comment_student='Зробила додаткове сортування за спаданням.',
    submitted_at=now - datetime.timedelta(hours=1)
)

sub3 = Submission.objects.create(
    assignment=task_multiclass,
    first_name='Денис',
    last_name='Кравченко',
    class_group=class_objs['8А'],
    teacher=teacher,
    comment_student='Трохи затримав здачу, вибачте!',
    submitted_at=now - datetime.timedelta(minutes=30)
)

# Робота, яка вже оцінена (для перевірки журналу)
sub_graded = Submission.objects.create(
    assignment=task_multiclass,
    first_name='Олена',
    last_name='Ткаченко',
    class_group=class_objs['9А'],
    teacher=teacher,
    grade='12',
    graded_at=now,
    comment_student='Робота виконана достроково.'
)
SubmissionComment.objects.create(
    submission=sub_graded,
    author=admin_user,
    text='Відмінно! Код чистий та відповідає стандарту PEP 8.'
)

print("\n" + "="*60)
print("🎉 ТЕСТОВЕ СЕРЕДОВИЩЕ УСПІШНО НАЛАШТОВАНО!")
print("="*60)
print(f"👤 Логін адміністратора/вчителя : {ADMIN_USERNAME}")
print(f"🔑 Пароль                        : {ADMIN_PASSWORD}")
print(f"👨‍🏫 ПІБ вчителя                  : {ADMIN_FULL_NAME}")
print("="*60)
print(f"📊 Створено предметів           : {Subject.objects.count()}")
print(f"🏫 Створено класів              : {ClassGroup.objects.count()}")
print(f"⏰ Створено дзвінків            : {BellSchedule.objects.count()}")
print(f"📅 Створено уроків у розкладі   : {TeacherLessonSchedule.objects.count()}")
print(f"📝 Створено завдань             : {Assignment.objects.count()}")
print(f"📥 Створено учнівських робіт    : {Submission.objects.count()} (з них неоцінених: {teacher.pending_reviews_count})")
print("="*60)
